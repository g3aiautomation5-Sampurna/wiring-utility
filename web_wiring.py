import os
import sys

# Ensure working directory is always the application folder (supports both PyInstaller .exe and script)
if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
    os.chdir(BASE_DIR)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    os.chdir(BASE_DIR)

import io
import base64
import shutil
import csv
from datetime import datetime
import cv2
import numpy as np
import time
import json
from flask import Flask, render_template_string, request, jsonify, send_file
import webbrowser
import threading

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pms_mqtt import pms_mqtt

from detect_origin import check_origin
from detect_origin2 import check_origin3_grommet, check_origin3_image, check_origin2_grommet, check_origin2_image
from ultralytics import YOLO

from pms_json_manager import (
    get_pc_mac_address,
    parse_sticker,
    build_device_autowhitelisting,
    build_serial_no_validation,
    build_status_validation,
    build_report_store,
    build_empty_report_store,
    build_empty_payloads,
    save_payload_to_file,
    load_or_create_config,
    update_config_value,
    CONFIG_FILE,
    validate_config_json
)

app = Flask(__name__)

# ================= VERSION MANAGEMENT =================
# APP_VERSION = "v1"
# ================= VERSION MANAGEMENT =================
APP_VERSION = str(load_or_create_config().get("slaveVersionNo", "") or "").strip()

# Change this to your own admin password.
ADMIN_VERSION_PASSWORD = "Admin@123"

# Ensure Result directories exist
os.makedirs(os.path.join(os.getcwd(), "Result", "Pass"), exist_ok=True)
os.makedirs(os.path.join(os.getcwd(), "Result", "Fail"), exist_ok=True)

EXCEL_PATH = os.path.join(os.getcwd(), "Result", "Inspection_Log.xlsx")
CSV_PATH = os.path.join(os.getcwd(), "Result", "Inspection_Log.csv")

def log_to_excel(record):
    """
    Appends the inspection record to Result/Inspection_Log.xlsx and Result/Inspection_Log.csv.
    """
    os.makedirs(os.path.join(os.getcwd(), "Result"), exist_ok=True)

    headers = [
        "Timestamp",
        "Device / Serial No.",
        "Overall Verdict",
        "Origin 1 (BL)",
        "Origin 2 (UM Coil)",
        "Origin 3 (Grommet)",
        "Origin 3 (BR Top)",
        "Wire Case 1 (BL)",
        "Wire Case 2 (UM)",
        "Wire Case 3 (BR)",
        "Diagnostic Summary",
        "Snapshot Folder"
    ]

    row_data = [
        record.get("timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        record.get("device_id", "N/A"),
        "PASS" if record.get("combined_pass") else "FAIL",
        "PASS" if record.get("obl_pass") else "FAIL",
        "PASS" if record.get("oum_pass") else "FAIL",
        "PASS" if record.get("orig3_grommet_pass") else "FAIL",
        "PASS" if record.get("obr_pass") else "FAIL",
        "PASS" if record.get("c1_pass") else "FAIL",
        "PASS" if record.get("c2_pass") else "FAIL",
        "PASS" if record.get("c3_pass") else "FAIL",
        record.get("notes_summary", ""),
        record.get("saved_folder", "")
    ]

    # 1. Update CSV Log
    csv_exists = os.path.exists(CSV_PATH)
    try:
        with open(CSV_PATH, "a", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            if not csv_exists:
                writer.writerow(headers)
            writer.writerow(row_data)
    except Exception as e:
        print(f"Warning: could not write to CSV log: {e}")

    # 2. Update Excel (.xlsx) Log with styling
    file_exists = os.path.exists(EXCEL_PATH)
    if file_exists:
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH)
            ws = wb.active
        except Exception:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Inspection Log"
            file_exists = False
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inspection Log"

    if not file_exists:
        ws.append(headers)
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        for col_num, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=col_num)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 26

    ws.append(row_data)
    current_row = ws.max_row
    ws.row_dimensions[current_row].height = 22

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    pass_font = Font(name="Segoe UI", size=10, bold=True, color="166534")

    fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fail_font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

    regular_font = Font(name="Segoe UI", size=10, color="0F172A")

    for col_idx in range(1, len(row_data) + 1):
        cell = ws.cell(row=current_row, column=col_idx)
        cell.border = thin_border
        cell.font = regular_font
        cell.alignment = Alignment(vertical="center", horizontal="center" if col_idx <= 10 else "left")

        val = str(cell.value)
        if val == "PASS":
            cell.fill = pass_fill
            cell.font = pass_font
        elif val == "FAIL":
            cell.fill = fail_fill
            cell.font = fail_font

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    try:
        wb.save(EXCEL_PATH)
    except Exception as e:
        print(f"Warning: could not save Excel file (it may be open): {e}")

def draw_labeled_box(img, pt1, pt2, is_pass, label):
    """Draws bounding box with clear colored badge banner."""
    x1, y1 = pt1
    x2, y2 = pt2
    border_color = (0, 220, 0) if is_pass else (0, 0, 240)
    cv2.rectangle(img, (x1, y1), (x2, y2), border_color, 3)

    badge_text = f" {label}: {'PASS' if is_pass else 'FAIL'} "
    badge_bg = (0, 180, 0) if is_pass else (0, 0, 220)
    
    badge_y = y1 - 8 if y1 > 25 else y1 + 24
    text_size = cv2.getTextSize(badge_text, cv2.FONT_HERSHEY_SIMPLEX, 0.52, 2)[0]
    cv2.rectangle(img, (x1, badge_y - 18), (x1 + text_size[0], badge_y + 4), badge_bg, -1)
    cv2.putText(img, badge_text, (x1, badge_y), cv2.FONT_HERSHEY_SIMPLEX, 0.52, (255, 255, 255), 2)

def decode_image(data_uri_or_bytes):
    if isinstance(data_uri_or_bytes, str) and data_uri_or_bytes.startswith("data:image"):
        header, encoded = data_uri_or_bytes.split(",", 1)
        data = base64.b64decode(encoded)
        nparr = np.frombuffer(data, np.uint8)
        return cv2.imdecode(nparr, cv2.IMREAD_COLOR)
    elif isinstance(data_uri_or_bytes, bytes):
        nparr = np.frombuffer(data_uri_or_bytes, np.uint8)
        return cv2.imdecode(nparr, cv2.IMREAD_COLOR)
    return None

def encode_image_to_base64(img_bgr):
    ret, buf = cv2.imencode(".jpg", img_bgr)
    if not ret:
        return None
    return "data:image/jpeg;base64," + base64.b64encode(buf).decode("utf-8")

# ================= YOLO MODELS LOADING =================
TOP_POSE_MODEL_PATH = os.path.join(BASE_DIR, "models", "wire_pose_model.pt")
if not os.path.exists(TOP_POSE_MODEL_PATH):
    TOP_POSE_MODEL_PATH = os.path.join(BASE_DIR, "models", "yolov8n-pose.pt")

SIDE_MODEL_PATH = os.path.join(BASE_DIR, "models", "origin3_model.pt")
if not os.path.exists(SIDE_MODEL_PATH):
    SIDE_MODEL_PATH = os.path.join(BASE_DIR, "models", "yolov8n.pt")

if not os.path.exists(TOP_POSE_MODEL_PATH) and os.path.isdir(os.path.join(BASE_DIR, "models")):
    for _fname in os.listdir(os.path.join(BASE_DIR, "models")):
        if _fname.endswith(".pt") and "pose" in _fname.lower():
            TOP_POSE_MODEL_PATH = os.path.join(BASE_DIR, "models", _fname)
            break

if not os.path.exists(SIDE_MODEL_PATH) and os.path.isdir(os.path.join(BASE_DIR, "models")):
    for _fname in os.listdir(os.path.join(BASE_DIR, "models")):
        if _fname.endswith(".pt") and os.path.join(BASE_DIR, "models", _fname) != TOP_POSE_MODEL_PATH:
            SIDE_MODEL_PATH = os.path.join(BASE_DIR, "models", _fname)
            break

top_pose_model = None
side_model = None

try:
    if os.path.exists(TOP_POSE_MODEL_PATH):
        top_pose_model = YOLO(TOP_POSE_MODEL_PATH)
        print(f"[YOLO] Loaded top pose estimation model from {TOP_POSE_MODEL_PATH}")
    else:
        print(f"[YOLO Warning] Top pose model file not found at {TOP_POSE_MODEL_PATH}")
except Exception as e:
    print(f"[YOLO Error] Failed to load top pose model: {e}")

try:
    if os.path.exists(SIDE_MODEL_PATH):
        side_model = YOLO(SIDE_MODEL_PATH)
        print(f"[YOLO] Loaded side detection model from {SIDE_MODEL_PATH}")
    else:
        print(f"[YOLO Warning] Side model file not found at {SIDE_MODEL_PATH}")
except Exception as e:
    print(f"[YOLO Error] Failed to load side model: {e}")


def evaluate_top_pose_model(model, img_bgr):
    """
    Infers on top image with the wire pose estimation model.
    Evaluates:
      - Bottom cases (Bottom Left & Bottom Right):
        Red wire's socket or terminal should be on the left to the socket or terminal of black wire.
      - Upper middle:
        Red wire's socket or terminal should be on the right of the socket or terminal of black wire.

      - Below left wire origins:
        Red wire's origin should be below black wire's origin.
      - Below right wire origins:
        Both wire origins should just exist.
      - Upper middle wire origins:
        The red wire's origin should be on the left of the black wire's origin.

    When model runs and results into Pass -> do nothing
    When it fails -> in an fstring store what all cases have failed
    """
    if img_bgr is None:
        failed_cases_str = f"Failed Cases: Case 1 (BL), Case 2 (UM), Case 3 (BR) (Image missing)"
        return False, None, [], False, False, False, False, False, False, failed_cases_str, []

    img_resized = cv2.resize(img_bgr, (1024, 576))
    if model is None:
        failed_cases_str = f"Failed Cases: Case 1 (BL), Case 2 (UM), Case 3 (BR) (Model not loaded)"
        return False, img_resized, ["Model not loaded"], False, False, False, False, False, False, failed_cases_str, []

    results = model(img_resized, verbose=False)
    res = results[0]
    annotated = res.plot()

    detected_labels = []
    names = res.names or {}
    kpts_data = res.keypoints.data.cpu().numpy() if res.keypoints is not None and len(res.keypoints.data) > 0 else None

    c1_red = c1_blk = c2_red = c2_blk = c3_red = c3_blk = None
    has_defect = False

    if len(res.boxes) > 0:
        for idx, box in enumerate(res.boxes):
            cls_id = int(box.cls[0].item()) if hasattr(box.cls, '__len__') else int(box.cls)
            conf = float(box.conf[0].item()) if hasattr(box.conf, '__len__') else float(box.conf)
            name = str(names.get(cls_id, cls_id)).lower()

            # Ignore background COCO objects
            if any(w in name for w in ['person', 'suitcase', 'backpack', 'tie', 'bottle', 'chair', 'cup', 'bed', 'laptop']):
                continue

            detected_labels.append(f"{name} ({conf:.2f})")
            if any(w in name for w in ['fail', 'defect', 'ng', 'bad', 'wrong', 'error']):
                has_defect = True

            if 'wire_red' not in name and 'wire_black' not in name:
                continue

            xyxy = box.xyxy[0].tolist() if hasattr(box.xyxy, '__len__') else box.xyxy.tolist()
            cx = (xyxy[0] + xyxy[2]) / (2.0 * 1024.0)
            cy = (xyxy[1] + xyxy[3]) / (2.0 * 576.0)
            kp = kpts_data[idx] if kpts_data is not None and idx < len(kpts_data) else None

            wire_info = {'name': name, 'conf': conf, 'kp': kp, 'xyxy': xyxy}

            if cy < 0.40:
                # Upper Middle (Case 2 / Origin 2)
                if 'wire_red' in name and (c2_red is None or conf > c2_red['conf']):
                    c2_red = wire_info
                elif 'wire_black' in name and (c2_blk is None or conf > c2_blk['conf']):
                    c2_blk = wire_info
            elif cx < 0.50:
                # Below Left (Case 1 / Origin 1)
                if 'wire_red' in name and (c1_red is None or conf > c1_red['conf']):
                    c1_red = wire_info
                elif 'wire_black' in name and (c1_blk is None or conf > c1_blk['conf']):
                    c1_blk = wire_info
            else:
                # Below Right (Case 3 / Origin 3 Top)
                if 'wire_red' in name and (c3_red is None or conf > c3_red['conf']):
                    c3_red = wire_info
                elif 'wire_black' in name and (c3_blk is None or conf > c3_blk['conf']):
                    c3_blk = wire_info

    top_fail_details = []

    # 1. Below Left (Case 1 / Origin 1)
    # Origin rule: Red wire's origin should be below black wire's origin (red.y > blk.y in image coords)
    # Terminal rule: Red wire's socket or terminal should be on the left of black wire's socket or terminal (red.x < blk.x)
    if c1_red is None or c1_blk is None or c1_red['kp'] is None or c1_blk['kp'] is None:
        obl_pass = False
        c1_pass = False
        missing = []
        if c1_red is None: missing.append("red wire")
        if c1_blk is None: missing.append("black wire")
        missing_str = ", ".join(missing) if missing else "wires"
        top_fail_details.append({"item": "Case 1: Bottom Left", "reason": f"Missing {missing_str}"})
        top_fail_details.append({"item": "Origin 1 (Bottom Left)", "reason": f"Missing {missing_str}"})
    else:
        # kp 0 = origin, kp 2 = terminal
        obl_pass = bool(c1_red['kp'][0][1] > c1_blk['kp'][0][1] and not has_defect)
        c1_term_pass = bool(c1_red['kp'][2][0] < c1_blk['kp'][2][0] and not has_defect)
        c1_pass = bool(c1_term_pass and not has_defect)
        if not obl_pass:
            top_fail_details.append({"item": "Origin 1 (Bottom Left)", "reason": "Red wire origin is not below black wire origin"})
        if not c1_pass:
            top_fail_details.append({"item": "Case 1: Bottom Left", "reason": "Red wire terminal is not to the left of black wire terminal"})

    # 2. Upper Middle (Case 2 / Origin 2)
    # Origin rule: The red wire's origin should be on the left of the black wire's origin (red.x < blk.x)
    # Terminal rule: The red wire's socket or terminal should be on the right of the socket or terminal of black wire (red.x > blk.x)
    if c2_red is None or c2_blk is None or c2_red['kp'] is None or c2_blk['kp'] is None:
        oum_pass = False
        c2_pass = False
        missing = []
        if c2_red is None: missing.append("red wire")
        if c2_blk is None: missing.append("black wire")
        missing_str = ", ".join(missing) if missing else "wires"
        top_fail_details.append({"item": "Case 2: Upper Middle", "reason": f"Missing {missing_str}"})
        top_fail_details.append({"item": "Origin 2 (Upper Middle)", "reason": f"Missing {missing_str}"})
    else:
        oum_pass = bool(c2_red['kp'][0][0] < c2_blk['kp'][0][0] and not has_defect)
        c2_term_pass = bool(c2_red['kp'][2][0] > c2_blk['kp'][2][0] and not has_defect)
        c2_pass = bool(c2_term_pass and not has_defect)
        if not oum_pass:
            top_fail_details.append({"item": "Origin 2 (Upper Middle)", "reason": "Red wire origin is not to the left of black wire origin"})
        if not c2_pass:
            top_fail_details.append({"item": "Case 2: Upper Middle", "reason": "Red wire terminal is not to the right of black wire terminal"})

    # 3. Below Right (Case 3 / Origin 3 Top)
    # Origin rule: Both wire origins should just exist
    # Terminal rule: The red wire's socket or terminal should be on the left to the socket or terminal of the black wire (red.x < blk.x)
    if c3_red is None or c3_blk is None or c3_red['kp'] is None or c3_blk['kp'] is None:
        obr_pass = False
        c3_pass = False
        missing = []
        if c3_red is None: missing.append("red wire")
        if c3_blk is None: missing.append("black wire")
        missing_str = ", ".join(missing) if missing else "wires"
        top_fail_details.append({"item": "Case 3: Bottom Right", "reason": f"Missing {missing_str}"})
        top_fail_details.append({"item": "Origin 3 (Bottom Right)", "reason": f"Missing {missing_str}"})
    else:
        obr_pass = bool(c3_red['kp'][0][2] > 0.20 and c3_blk['kp'][0][2] > 0.20 and not has_defect)
        c3_term_pass = bool(c3_red['kp'][2][0] < c3_blk['kp'][2][0] and not has_defect)
        c3_pass = bool(c3_term_pass and not has_defect)
        if not obr_pass:
            top_fail_details.append({"item": "Origin 3 (Bottom Right)", "reason": "Wire origins not detected"})
        if not c3_term_pass:
            top_fail_details.append({"item": "Case 3: Bottom Right", "reason": "Red wire terminal is not to the left of black wire terminal"})

    is_top_pass = bool(obl_pass and oum_pass and obr_pass and c1_pass and c2_pass and c3_pass and not has_defect)

    # When the model runs and results into Pass then do nothing
    # and if it fails then in an fstring store what all cases have failed
    if is_top_pass:
        failed_cases_str = ""
    else:
        failed_cases = []
        if not obl_pass: failed_cases.append("Origin 1 (BL)")
        if not oum_pass: failed_cases.append("Origin 2 (UM)")
        if not obr_pass: failed_cases.append("Origin 3 (BR Top)")
        if not c1_pass: failed_cases.append("Case 1 (BL)")
        if not c2_pass: failed_cases.append("Case 2 (UM)")
        if not c3_pass: failed_cases.append("Case 3 (BR)")
        failed_list = ", ".join(failed_cases)
        failed_cases_str = f"Failed Cases: {failed_list}"

    return (
        is_top_pass, annotated, detected_labels,
        obl_pass, oum_pass, obr_pass,
        c1_pass, c2_pass, c3_pass,
        failed_cases_str, top_fail_details
    )


def evaluate_side_model(model, img_bgr):
    """
    Infers on side image with the Origin 3 detection model.
    When model runs and results into Pass -> do nothing
    When it fails -> in an fstring store what all cases have failed
    """
    if img_bgr is None:
        failed_cases_str = f"Failed Cases: Origin 3 (Side View) (Image missing)"
        return False, None, [], failed_cases_str

    img_resized = cv2.resize(img_bgr, (1024, 576))
    if model is None:
        failed_cases_str = f"Failed Cases: Origin 3 (Side View) (Model not loaded)"
        return False, img_resized, ["Model not loaded"], failed_cases_str

    results = model(img_resized, verbose=False)
    res = results[0]
    annotated = res.plot()

    detected_labels = []
    has_defect = False
    has_origin3 = False
    names = res.names or {}

    if len(res.boxes) > 0:
        for box in res.boxes:
            cls_id = int(box.cls[0].item()) if hasattr(box.cls, '__len__') else int(box.cls)
            conf = float(box.conf[0].item()) if hasattr(box.conf, '__len__') else float(box.conf)
            name = str(names.get(cls_id, cls_id)).lower()

            if any(w in name for w in ['person', 'suitcase', 'backpack', 'tie', 'bottle', 'chair', 'cup', 'bed', 'laptop']):
                continue

            detected_labels.append(f"{name} ({conf:.2f})")
            if any(w in name for w in ['fail', 'defect', 'ng', 'bad', 'wrong', 'error']):
                has_defect = True
            if 'origin_3' in name or 'origin3' in name or 'pass' in name or 'ok' in name:
                has_origin3 = True

    is_side_pass = bool(has_origin3 and not has_defect)

    # When the model runs and results into Pass then do nothing
    # and if it fails then in an fstring store what all cases have failed
    if is_side_pass:
        failed_cases_str = ""
    else:
        failed_cases_str = f"Failed Cases: Origin 3 (Side View)"

    return is_side_pass, annotated, detected_labels, failed_cases_str


def evaluate_yolo_detection(model, img_bgr, is_pose=False):
    """Compatibility helper for generic YOLO evaluation."""
    if is_pose:
        (
            is_pass, annotated, detected_labels,
            obl_pass, oum_pass, obr_pass,
            c1_pass, c2_pass, c3_pass,
            _, _
        ) = evaluate_top_pose_model(model, img_bgr)
    else:
        is_pass, annotated, detected_labels, _ = evaluate_side_model(model, img_bgr)
    return is_pass, annotated, detected_labels


def process_top_inspection(img_raw):
    """
    Top image detected by the pose estimation model.
    Evaluates:
      - Wire Origins (obl_pass, oum_pass, obr_pass)
      - Wire Routing (c1_pass, c2_pass, c3_pass)
    """
    (
        is_top_pass, img_annotated, detected_labels,
        obl_pass, oum_pass, obr_pass,
        c1_pass, c2_pass, c3_pass,
        failed_cases_str, top_fail_details
    ) = evaluate_top_pose_model(top_pose_model, img_raw)

    if img_annotated is None:
        img_annotated = cv2.resize(img_raw, (1024, 576)) if img_raw is not None else np.zeros((576, 1024, 3), dtype=np.uint8)

    encoded_annotated = encode_image_to_base64(img_annotated)

    return {
        "top_wires_annotated": encoded_annotated,
        "top_origins_annotated": encoded_annotated,
        "img_wires": img_annotated,
        "img_origins": img_annotated,
        "top_pass": bool(is_top_pass),
        "obl_pass": bool(obl_pass),
        "oum_pass": bool(oum_pass),
        "obr_pass": bool(obr_pass),
        "c1_pass": bool(c1_pass),
        "c2_pass": bool(c2_pass),
        "c3_pass": bool(c3_pass),
        "detected_labels": detected_labels,
        "failed_cases_str": failed_cases_str,
        "fail_details": top_fail_details
    }


def process_orig3_inspection(img_raw):
    """
    Side image detected by the other model.
    Evaluates Origin 3 (Side View).
    """
    is_side_pass, img_annotated, detected_labels, failed_cases_str = evaluate_side_model(side_model, img_raw)
    if img_annotated is None:
        img_annotated = cv2.resize(img_raw, (1024, 576)) if img_raw is not None else np.zeros((576, 1024, 3), dtype=np.uint8)

    encoded_annotated = encode_image_to_base64(img_annotated)
    failed_reasons = [] if is_side_pass else ["Side model did not detect Origin 3."]

    return {
        "orig3_annotated": encoded_annotated,
        "orig2_annotated": encoded_annotated,
        "img_o3": img_annotated,
        "orig3_grommet_pass": bool(is_side_pass),
        "orig2_grommet_pass": bool(is_side_pass),
        "side_pass": bool(is_side_pass),
        "failed_reasons": failed_reasons,
        "detected_labels": detected_labels,
        "failed_cases_str": failed_cases_str
    }

process_orig2_inspection = process_orig3_inspection

HTML_PAGE = """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Wiring & Origin Inspection</title>
    <style>
        :root {
            --bg-base: #0b0f19;
            --bg-surface: #121826;
            --bg-surface-elevated: #1a2333;
            --border-subtle: #1e293b;
            --border-highlight: #334155;
            --text-main: #f1f5f9;
            --text-secondary: #94a3b8;
            --text-muted: #64748b;
            --accent-primary: #3b82f6;
            --accent-primary-hover: #2563eb;
            --color-pass: #10b981;
            --color-pass-bg: rgba(16, 185, 129, 0.12);
            --color-fail: #ef4444;
            --color-fail-bg: rgba(239, 68, 68, 0.12);
            --radius-sm: 6px;
            --radius-md: 10px;
            --radius-lg: 14px;
        }

        * {
            box-sizing: border-box;
            margin: 0;
            padding: 0;
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
        }

        body {
            background-color: var(--bg-base);
            color: var(--text-main);
            min-height: 100vh;
            display: flex;
            flex-direction: column;
            overflow-x: hidden;
            font-size: 14px;
        }

        /* Minimal Header */
        header {
            background: var(--bg-surface);
            border-bottom: 1px solid var(--border-subtle);
            padding: 8px 18px;
            display: flex;
            align-items: center;
            justify-content: space-between;
            gap: 14px;
            min-height: 56px;
            flex-wrap: wrap;
        }

        .brand-group {
            display: flex;
            align-items: center;
            gap: 10px;
        }

        .brand-icon {
            width: 28px;
            height: 28px;
            background: linear-gradient(135deg, #3b82f6, #6366f1);
            border-radius: var(--radius-sm);
            display: flex;
            align-items: center;
            justify-content: center;
            font-weight: 800;
            font-size: 14px;
            color: #fff;
        }

        .brand-title {
            font-size: 15px;
            font-weight: 700;
            letter-spacing: -0.2px;
            color: #ffffff;
        }

        
        .brand-tag {
            font-size: 14px;
            font-weight: 700;
            color: #ffffff;
            background: #24334a;
            padding: 5px 11px;
            border-radius: 5px;
            border: 1px solid #526985;
            letter-spacing: 0.3px;
        }


        /* Pill Navigation in Header */
        .nav-pills {
            display: flex;
            background: var(--bg-base);
            border: 1px solid var(--border-subtle);
            border-radius: var(--radius-sm);
            padding: 3px;
            gap: 2px;
        }

        .nav-pill {
            background: transparent;
            border: none;
            color: var(--text-secondary);
            font-size: 12px;
            font-weight: 600;
            padding: 5px 12px;
            border-radius: 4px;
            cursor: pointer;
            transition: all 0.15s ease;
            display: inline-flex;
            align-items: center;
            gap: 5px;
        }

        .nav-pill:hover {
            color: var(--text-main);
        }

        .nav-pill.active {
            background: var(--bg-surface-elevated);
            color: #60a5fa;
            box-shadow: 0 1px 3px rgba(0,0,0,0.3);
        }

        .nav-divider {
            width: 1px;
            height: 18px;
            background: var(--border-subtle);
            margin: auto 4px;
        }

        /* Header Right Controls */
        .header-actions {
            display: flex;
            align-items: center;
            gap: 8px;
            flex-wrap: wrap;
        }

        /* QR Code Input Container */
        .qr-input-group {
            display: flex;
            align-items: center;
            background: var(--bg-base);
            border: 1px solid var(--border-highlight);
            border-radius: var(--radius-sm);
            padding: 3px 8px;
            gap: 6px;
            transition: all 0.2s ease;
        }

        .qr-input-group:focus-within {
            border-color: #3b82f6;
            box-shadow: 0 0 0 2px rgba(59, 130, 246, 0.2);
        }

        .qr-icon {
            font-size: 13px;
            opacity: 0.85;
        }

        .qr-input {
            background: transparent;
            border: none;
            color: #f8fafc;
            font-size: 12px;
            font-weight: 600;
            width: 210px;
            outline: none;
        }

        .qr-badge {
            font-size: 10px;
            padding: 2px 7px;
            border-radius: 4px;
            font-weight: 700;
            white-space: nowrap;
            user-select: none;
            letter-spacing: 0.3px;
            transition: all 0.2s ease;
        }

        .qr-badge.idle {
            background: rgba(100, 116, 139, 0.2);
            color: #94a3b8;
            border: 1px solid #334155;
        }

        .qr-badge.validating {
            background: rgba(245, 158, 11, 0.2);
            color: #fbbf24;
            border: 1px solid #f59e0b;
        }

        .qr-badge.success {
            background: rgba(16, 185, 129, 0.25);
            color: #34d399;
            border: 1px solid #10b981;
        }

        .qr-badge.failed {
            background: rgba(239, 68, 68, 0.25);
            color: #f87171;
            border: 1px solid #ef4444;
        }

        .qr-input::placeholder {
            color: var(--text-muted);
            font-weight: 400;
        }

        .btn-clear-qr {
            background: transparent;
            border: none;
            color: var(--text-muted);
            font-size: 11px;
            cursor: pointer;
            padding: 0 2px;
            line-height: 1;
        }

        .btn-clear-qr:hover {
            color: #f87171;
        }

        .btn-header {
            background: var(--bg-surface-elevated);
            color: var(--text-secondary);
            border: 1px solid var(--border-subtle);
            padding: 6px 12px;
            border-radius: var(--radius-sm);
            font-size: 12px;
            font-weight: 600;
            cursor: pointer;
            display: inline-flex;
            align-items: center;
            gap: 6px;
            transition: all 0.15s ease;
        }

        .btn-header:hover {
            color: var(--text-main);
            border-color: var(--border-highlight);
            background: #222d42;
        }

        /* PMS JSON Modal */
        .modal-backdrop {
            position: fixed;
            top: 0; left: 0; right: 0; bottom: 0;
            background: rgba(0, 0, 0, 0.75);
            backdrop-filter: blur(4px);
            display: flex;
            align-items: center;
            justify-content: center;
            z-index: 9999;
        }
        .modal-box {
            background: #111827;
            border: 1px solid #374151;
            border-radius: 8px;
            width: 760px;
            max-width: 92vw;
            max-height: 85vh;
            display: flex;
            flex-direction: column;
            box-shadow: 0 20px 40px rgba(0,0,0,0.6);
        }
        .modal-header {
            display: flex;
            align-items: center;
            justify-content: space-between;
            padding: 14px 20px;
            border-bottom: 1px solid #374151;
        }
        .modal-header h3 {
            margin: 0;
            font-size: 15px;
            font-weight: 600;
            color: #f3f4f6;
        }
        .modal-close {
            background: none;
            border: none;
            color: #9ca3af;
            font-size: 22px;
            cursor: pointer;
            line-height: 1;
        }
        .modal-close:hover { color: #fff; }
        .modal-tabs {
            display: flex;
            gap: 8px;
            padding: 10px 20px;
            background: #162032;
            border-bottom: 1px solid #374151;
            overflow-x: auto;
        }
        .modal-tab {
            background: transparent;
            border: 1px solid transparent;
            color: #94a3b8;
            font-size: 11px;
            font-weight: 600;
            padding: 6px 12px;
            border-radius: 4px;
            cursor: pointer;
            white-space: nowrap;
        }
        .modal-tab.active {
            background: #3b82f6;
            color: #fff;
        }
        .modal-pre {
            flex: 1;
            margin: 0;
            padding: 16px 20px;
            background: #0b0f19;
            color: #38bdf8;
            font-family: 'Consolas', 'Courier New', monospace;
            font-size: 12px;
            line-height: 1.5;
            overflow: auto;
            max-height: 52vh;
        }
        .modal-footer {
            display: flex;
            justify-content: flex-end;
            gap: 10px;
            padding: 12px 20px;
            border-top: 1px solid #374151;
            background: #162032;
        }
        .btn-copy-json {
            background: #3b82f6;
            color: #fff;
            border: none;
            padding: 6px 14px;
            border-radius: 4px;
            font-size: 12px;
            font-weight: 600;
            cursor: pointer;
        }
        .btn-copy-json:hover { background: #2563eb; }
        .btn-modal-close {
            background: #1f2937;
            color: #d1d5db;
            border: 1px solid #374151;
            padding: 6px 14px;
            border-radius: 4px;
            font-size: 12px;
            cursor: pointer;
        }
        .btn-modal-close:hover { color: #fff; }

        /* Slide In/Out Config Toast Popup */
        #configToastContainer {
            position: fixed;
            top: 24px;
            right: 24px;
            z-index: 99999;
            pointer-events: none;
            display: flex;
            flex-direction: column;
            gap: 12px;
        }

        .config-toast {
            pointer-events: auto;
            min-width: 360px;
            max-width: 460px;
            background: #111827;
            border: 1px solid #374151;
            border-radius: 8px;
            box-shadow: 0 12px 30px rgba(0, 0, 0, 0.6), 0 2px 6px rgba(0, 0, 0, 0.4);
            overflow: hidden;
            position: relative;
            transform: translateX(120%);
            opacity: 0;
            transition: transform 0.35s cubic-bezier(0.16, 1, 0.3, 1), opacity 0.35s ease;
            cursor: pointer;
        }

        .config-toast.show {
            transform: translateX(0);
            opacity: 1;
        }

        .config-toast.hide {
            transform: translateX(120%);
            opacity: 0;
        }

        .config-toast.success {
            border-left: 5px solid #10b981;
        }

        .config-toast.failed {
            border-left: 5px solid #ef4444;
        }

        .toast-content {
            display: flex;
            align-items: flex-start;
            gap: 12px;
            padding: 14px 16px;
        }

        .toast-icon {
            width: 32px;
            height: 32px;
            border-radius: 50%;
            display: flex;
            align-items: center;
            justify-content: center;
            font-size: 16px;
            font-weight: bold;
            flex-shrink: 0;
            margin-top: 1px;
        }

        .toast-icon.success {
            background: rgba(16, 185, 129, 0.2);
            color: #10b981;
            border: 1px solid rgba(16, 185, 129, 0.4);
        }

        .toast-icon.failed {
            background: rgba(239, 68, 68, 0.2);
            color: #ef4444;
            border: 1px solid rgba(239, 68, 68, 0.4);
        }

        .toast-text {
            flex: 1;
            display: flex;
            flex-direction: column;
            gap: 4px;
        }

        .toast-header-row {
            display: flex;
            align-items: center;
            justify-content: space-between;
            gap: 8px;
        }

        .toast-title {
            font-size: 14px;
            font-weight: 700;
            color: #f8fafc;
        }

        .toast-badge {
            font-size: 10px;
            font-weight: 700;
            letter-spacing: 0.5px;
            padding: 2px 7px;
            border-radius: 3px;
        }

        .toast-badge.success {
            background: rgba(16, 185, 129, 0.2);
            color: #34d399;
            border: 1px solid #059669;
        }

        .toast-badge.failed {
            background: rgba(239, 68, 68, 0.2);
            color: #f87171;
            border: 1px solid #dc2626;
        }

        .toast-body {
            font-size: 12px;
            color: #cbd5e1;
            line-height: 1.45;
        }

        .toast-meta {
            display: flex;
            gap: 12px;
            font-size: 11px;
            color: #94a3b8;
            margin-top: 2px;
        }

        .toast-meta span strong {
            color: #38bdf8;
            font-family: 'Consolas', monospace;
        }

        .toast-close-btn {
            background: none;
            border: none;
            color: #94a3b8;
            font-size: 18px;
            cursor: pointer;
            padding: 0 4px;
            line-height: 1;
        }

        .toast-close-btn:hover {
            color: #fff;
        }

        /* Progress timer bar at the bottom */
        .toast-progress {
            position: absolute;
            bottom: 0;
            left: 0;
            height: 3px;
            width: 100%;
            background: rgba(255, 255, 255, 0.1);
        }

        .toast-progress-bar {
            height: 100%;
            width: 100%;
            transform-origin: left;
        }

        .toast-progress-bar.success {
            background: #10b981;
        }

        .toast-progress-bar.failed {
            background: #ef4444;
        }

        @keyframes toastProgressCountdown {
            from { transform: scaleX(1); }
            to { transform: scaleX(0); }
        }

        .btn-start {
            background: linear-gradient(135deg, #10b981, #059669);
            color: #ffffff;
            border: none;
            padding: 7px 18px;
            border-radius: var(--radius-sm);
            font-size: 13px;
            font-weight: 700;
            cursor: pointer;
            display: inline-flex;
            align-items: center;
            gap: 6px;
            box-shadow: 0 2px 8px rgba(16, 185, 129, 0.3);
            transition: all 0.2s ease;
        }

        .btn-start:not(:disabled):hover {
            background: linear-gradient(135deg, #059669, #047857);
            transform: translateY(-1px);
            box-shadow: 0 4px 14px rgba(16, 185, 129, 0.45);
        }

        .btn-start:not(:disabled):active {
            transform: translateY(0);
        }

        .btn-start:disabled {
            background: #1e293b;
            color: #64748b;
            border: 1px solid #334155;
            cursor: not-allowed;
            box-shadow: none;
            transform: none;
            opacity: 0.6;
        }

        .btn-start.ready-pulse {
            animation: pulse-ready 1.8s infinite;
        }

        @keyframes pulse-ready {
            0% { box-shadow: 0 0 0 0 rgba(16, 185, 129, 0.6); }
            70% { box-shadow: 0 0 0 8px rgba(16, 185, 129, 0); }
            100% { box-shadow: 0 0 0 0 rgba(16, 185, 129, 0); }
        }

        /* Layout */
        .main-layout {
            display: flex;
            flex: 1;
            padding: 14px 18px;
            gap: 14px;
            max-width: 1750px;
            margin: 0 auto;
            width: 100%;
        }

        .viewport-wrapper {
            flex: 1;
            display: flex;
            flex-direction: column;
            gap: 10px;
            min-width: 0;
        }

        .sidebar-wrapper {
            width: 360px;
            display: flex;
            flex-direction: column;
            gap: 10px;
            flex-shrink: 0;
        }

        /* Viewport Card */
        .viewport-card {
            background: var(--bg-surface);
            border: 1px solid var(--border-subtle);
            border-radius: var(--radius-md);
            overflow: hidden;
            display: flex;
            flex-direction: column;
            flex: 1;
            min-height: 520px;
        }

        .viewport-content {
            flex: 1;
            display: flex;
            position: relative;
            background: #06090e;
            align-items: center;
            justify-content: center;
            overflow: hidden;
        }

        /* Full Single Camera Viewport */
        .camera-full-container {
            width: 100%;
            height: 100%;
            position: absolute;
            top: 0;
            left: 0;
            display: flex;
            align-items: center;
            justify-content: center;
            cursor: pointer;
            transition: border-color 0.2s;
        }

        .camera-full-container.drag-over {
            background: rgba(59, 130, 246, 0.1);
        }

        .camera-overlay-badge {
            position: absolute;
            top: 14px;
            left: 14px;
            background: rgba(11, 15, 25, 0.85);
            backdrop-filter: blur(6px);
            border: 1px solid rgba(255, 255, 255, 0.1);
            padding: 5px 12px;
            border-radius: 6px;
            font-size: 12px;
            font-weight: 600;
            display: flex;
            align-items: center;
            gap: 8px;
            z-index: 10;
        }

        .badge-dot {
            width: 8px;
            height: 8px;
            border-radius: 50%;
            background: #64748b;
        }

        .badge-dot.active {
            background: var(--color-pass);
            box-shadow: 0 0 8px var(--color-pass);
        }

        .camera-actions-overlay {
            position: absolute;
            top: 14px;
            right: 14px;
            z-index: 10;
            display: flex;
            gap: 8px;
        }

        .btn-pill-action {
            background: rgba(18, 24, 38, 0.9);
            color: #cbd5e1;
            border: 1px solid rgba(255, 255, 255, 0.12);
            padding: 5px 10px;
            border-radius: 6px;
            font-size: 11px;
            font-weight: 600;
            cursor: pointer;
            backdrop-filter: blur(4px);
            transition: all 0.15s ease;
        }

        .btn-pill-action:hover {
            background: #3b82f6;
            color: #fff;
            border-color: #3b82f6;
        }

        .camera-full-container video, .camera-full-container img {
            max-width: 100%;
            max-height: 100%;
            object-fit: contain;
            user-select: none;
        }

        .guide-canvas {
            position: absolute;
            top: 0;
            left: 0;
            width: 100%;
            height: 100%;
            pointer-events: none;
        }

        /* Single Result View */
        .single-view {
            width: 100%;
            height: 100%;
            position: absolute;
            top: 0;
            left: 0;
            display: none;
            flex-direction: column;
            align-items: center;
            justify-content: center;
            padding: 8px;
        }

        .single-view img {
            max-width: 100%;
            max-height: 100%;
            object-fit: contain;
            border-radius: 4px;
        }

        /* Bottom Minimal Controls Bar */
        .bottom-bar {
            background: var(--bg-surface);
            border: 1px solid var(--border-subtle);
            border-radius: var(--radius-sm);
            padding: 7px 12px;
            display: flex;
            align-items: center;
            justify-content: space-between;
            font-size: 12px;
            color: var(--text-secondary);
        }

        .cam-controls {
            display: flex;
            align-items: center;
            gap: 8px;
        }

        .cam-controls select {
            background: var(--bg-surface-elevated);
            color: var(--text-main);
            border: 1px solid var(--border-subtle);
            border-radius: 4px;
            padding: 3px 8px;
            font-size: 11px;
            outline: none;
            max-width: 150px;
        }

        .btn-icon {
            background: var(--bg-surface-elevated);
            color: var(--text-secondary);
            border: 1px solid var(--border-subtle);
            padding: 3px 8px;
            border-radius: 4px;
            cursor: pointer;
            font-size: 11px;
        }

        .btn-icon:hover {
            color: var(--text-main);
            border-color: var(--border-highlight);
        }

        /* Sidebar Elements */
        .sidebar-card {
            background: var(--bg-surface);
            border: 1px solid var(--border-subtle);
            border-radius: var(--radius-md);
            padding: 14px;
        }

        /* Verdict Box */
       .verdict-box {
            display: flex;
            flex-direction: column;
            align-items: center;
            justify-content: center;
            gap: 8px;
            padding: 16px 20px;
            min-height: 110px;
            width: 100%;
            text-align: center;
            box-sizing: border-box;
        }

        .verdict-box.pass {
            background: var(--color-pass-bg);
            border-color: rgba(16, 185, 129, 0.4);
        }

        .verdict-box.fail {
            background: var(--color-fail-bg);
            border-color: rgba(239, 68, 68, 0.4);
        }

        .verdict-status {
            width: 100%;
            flex: 0 0 auto;
            font-size: 28px;
            font-weight: 800;
            line-height: 1.2;
            letter-spacing: 0.5px;
            color: var(--text-muted);
            display: flex;
            align-items: center;
            justify-content: center;
            text-align: center;
        }

        .verdict-box.pass .verdict-status {
            color: var(--color-pass);
        }

        .verdict-box.fail .verdict-status {
            color: var(--color-fail);
        }

        .verdict-device {
            font-size: 12px;
            font-weight: 700;
            color: var(--text-main);
            margin-top: 5px;
            display: none !important;
        }

        .verdict-desc {
            width: 100%;
            flex: 0 0 auto;
            min-width: 0;
            margin-top: 0;
            font-size: 14px;
            line-height: 1.5;
            color: var(--text-muted);
            text-align: center;
            white-space: normal;
            overflow-wrap: break-word;
            word-break: normal;
        }

        /* Checklist Section */
        .section-header {
            font-size: 11px;
            font-weight: 700;
            text-transform: uppercase;
            letter-spacing: 0.5px;
            color: var(--text-muted);
            margin-bottom: 8px;
            display: flex;
            align-items: center;
            justify-content: space-between;
        }

        .check-table {
            width: 100%;
            display: flex;
            flex-direction: column;
            gap: 3px;
        }

        .check-row {
            display: flex;
            align-items: center;
            justify-content: space-between;
            padding: 6px 8px;
            background: rgba(255, 255, 255, 0.02);
            border-radius: 4px;
            font-size: 12px;
        }

        .check-name {
            display: flex;
            align-items: center;
            gap: 8px;
            color: var(--text-main);
            font-weight: 500;
        }

        .status-pill {
            font-size: 11px;
            font-weight: 700;
            padding: 2px 7px;
            border-radius: 4px;
            background: #1e293b;
            color: #64748b;
        }

        .status-pill.pass {
            background: rgba(16, 185, 129, 0.15);
            color: #34d399;
        }

        .status-pill.fail {
            background: rgba(239, 68, 68, 0.15);
            color: #f87171;
        }

        /* Simplified Inspection Result Card */
        .result-idle-state {
            display: flex;
            align-items: center;
            padding: 14px 16px;
            background: rgba(255, 255, 255, 0.02);
            border: 1px dashed var(--border-subtle);
            border-radius: var(--radius-md);
            color: var(--text-muted);
        }

        .result-pass-state {
            display: flex;
            flex-direction: column;
            align-items: center;
            justify-content: center;
            padding: 24px 16px;
            background: rgba(16, 185, 129, 0.12);
            border: 1.5px solid rgba(16, 185, 129, 0.45);
            border-radius: var(--radius-md);
            text-align: center;
        }

        .pass-badge-large {
            font-size: 32px;
            font-weight: 800;
            letter-spacing: 2px;
            color: #10b981;
            text-shadow: 0 0 16px rgba(16, 185, 129, 0.35);
            display: flex;
            align-items: center;
            gap: 10px;
        }

        .result-fail-state {
            display: flex;
            flex-direction: column;
            padding: 16px 14px;
            background: rgba(239, 68, 68, 0.08);
            border: 1.5px solid rgba(239, 68, 68, 0.45);
            border-radius: var(--radius-md);
            gap: 12px;
        }

        .fail-badge-large {
            font-size: 26px;
            font-weight: 800;
            letter-spacing: 1.5px;
            color: #ef4444;
            text-shadow: 0 0 16px rgba(239, 68, 68, 0.35);
            display: flex;
            align-items: center;
            gap: 8px;
        }

        .fail-list-title {
            font-size: 11px;
            font-weight: 700;
            text-transform: uppercase;
            letter-spacing: 0.8px;
            color: #f87171;
            margin-top: 2px;
        }

        .fail-items-container {
            display: flex;
            flex-direction: column;
            gap: 8px;
        }

        .fail-item-row {
            display: flex;
            align-items: flex-start;
            gap: 10px;
            background: rgba(239, 68, 68, 0.12);
            border-left: 3px solid #ef4444;
            padding: 8px 10px;
            border-radius: 4px;
            font-size: 12px;
            color: #fee2e2;
            line-height: 1.4;
        }

        .fail-item-row .fail-icon {
            color: #ef4444;
            font-weight: bold;
            font-size: 13px;
            flex-shrink: 0;
            margin-top: 1px;
        }

        .fail-item-row .fail-text {
            flex: 1;
            font-weight: 500;
        }

        .fail-item-row .fail-name {
            font-weight: 700;
            color: #ffffff;
            font-size: 12px;
        }

        .fail-item-row .fail-subtext {
            font-size: 11px;
            color: #fca5a5;
            margin-top: 2px;
        }

        /* Diagnostics Accordion */
        details {
            background: var(--bg-surface);
            border: 1px solid var(--border-subtle);
            border-radius: var(--radius-md);
            overflow: hidden;
        }

        details summary {
            padding: 10px 14px;
            font-size: 12px;
            font-weight: 600;
            color: var(--text-secondary);
            cursor: pointer;
            user-select: none;
            outline: none;
        }

        details summary:hover {
            color: var(--text-main);
        }

        .diag-log {
            padding: 10px 14px;
            background: #080c14;
            color: #cbd5e1;
            font-family: 'Consolas', monospace;
            font-size: 11px;
            line-height: 1.4;
            white-space: pre-wrap;
            max-height: 140px;
            overflow-y: auto;
            border-top: 1px solid var(--border-subtle);
        }

        /* Pulse Flash on Start */
        @keyframes pulse-flash {
            0% { opacity: 0.3; }
            100% { opacity: 1; }
        }
        .flash-anim {
            animation: pulse-flash 0.25s ease-out;
        }

        input[type="file"] {
            display: none;
        }

        /* MQTT Connection Indicator */
        .mqtt-status {
            display: inline-flex;
            align-items: center;
            gap: 8px;
            padding: 7px 12px;
            border-radius: 20px;
            font-size: 12px;
            font-weight: 600;
            white-space: nowrap;
        }

        .mqtt-status-dot {
            width: 9px;
            height: 9px;
            border-radius: 50%;
            background: #ef4444;
            flex-shrink: 0;
        }

        .mqtt-status.disconnected {
            color: #fca5a5;
            background: rgba(239, 68, 68, 0.12);
        }

        .mqtt-status.connected {
            color: #86efac;
            background: rgba(34, 197, 94, 0.12);
        }

        .mqtt-status.connected .mqtt-status-dot {
            background: #22c55e;
            box-shadow: 0 0 7px rgba(34, 197, 94, 0.65);
        }
        .start-version-group {
            display: inline-flex;
            align-items: center;
            gap: 10px;
            flex-shrink: 0;
        }

        .version-display {
            font-size: 17px;
            font-weight: 700;
            color: #ffffff;
            background: #1f2937;
            padding: 6px 10px;
            border-radius: 6px;
            border: 1px solid #64748b;
            white-space: nowrap;
        }
    </style>
</head>
<body>
    <!-- HEADER -->
    <header>
        <div class="brand-group">
            <div class="brand-icon">W</div>
            <div>
                <span class="brand-title">Wiring Inspection</span>
                <span class="brand-tag">Dual Camera</span>
                <span class="brand-tag" id="appVersionBadge"></span>

            </div>
        </div>

        <!-- Pill Nav (Live feeds & Inspection Results) -->
        <div class="nav-pills">
            <button class="nav-pill active" id="pillLiveTop" onclick="setViewMode('live_top')">📹 Live Top View</button>
            <button class="nav-pill" id="pillLiveOrig" onclick="setViewMode('live_orig3')">📹 Live Origin 3</button>
            <div class="nav-divider"></div>
            <button class="nav-pill" id="pillTopWires" onclick="setViewMode('res_top_wires')">🔌 Top Wires</button>
            <button class="nav-pill" id="pillTopOrigins" onclick="setViewMode('res_top_origins')">📍 Top Origins</button>
            <button class="nav-pill" id="pillOrig3" onclick="setViewMode('res_orig3')">🎯 Origin 3</button>
        </div>

        <!-- Actions (QR input, Batch Upload, Excel, START) -->
        <div class="header-actions">
            <!-- QR / Barcode Scanner Input Box -->
            <div class="qr-input-group" title="Scan QR Code (Format: Id: ..; Sr: ..;) or Device Serial No.">
                <span class="qr-icon">🏷️</span>
                <input type="text" id="deviceSerialInput" class="qr-input" placeholder="Scan QR (Id: ..; Sr: ..;) or Serial..." autocomplete="off" autofocus />
                <span id="qrValidationBadge" class="qr-badge idle" title="PMS Validation Status">Scan to Validate</span>
                <button class="btn-clear-qr" id="btnClearQr" onclick="clearQrInput()" title="Clear Device ID">✕</button>
            </div>

            <button class="btn-header" onclick="document.getElementById('fileBothInput').click()" title="Select both test images">
                📁 Upload
            </button>
            <button class="btn-header" onclick="openExcelSheet()" title="Open Inspection Log in Microsoft Excel">
                📊 Excel Log
            </button>
            <button class="btn-header" onclick="openPmsJsonModal()" title="View & Copy Selec PMS JSON output">
                📋 PMS JSON
            </button>
            <button class="btn-header" onclick="openVersionAdminModal()" title="Admin version settings">
                ⚙️ Admin Settings
            </button>
            <div id="mqttStatus" class="mqtt-status disconnected">
                <span class="mqtt-status-dot"></span>
                <span id="mqttStatusText">SERVER  Disconnected</span>
            </div>
            <button class="btn-header" id="btnLiveReset" onclick="resetToLiveCameras()" title="Switch back to live video streams">
                🔄 Live
            </button>
            <div class="start-version-group">
                <button class="btn-start" id="btnStartInspection"
                        onclick="triggerStartInspection()" disabled
                        title="Scan QR Code and wait for PMS server validation (CONFIG SUCCESSFUL) to activate START">
                    ▶ START
                </button>

                <span class="version-display" id="startVersionBadge">GTI Version: 1</span>
            </div>
        </div>
    </header>

    <!-- MAIN BODY -->
    <div class="main-layout">
        <!-- LEFT: Viewport & Bottom Bar -->
        <div class="viewport-wrapper">
            <div class="viewport-card">
                <div class="viewport-content" id="viewportArea">
                    <!-- Live Top Camera Viewport (Full Screen) -->
                    <div class="camera-full-container" id="containerTop" ondragover="handleDragOver(event)" ondragleave="handleDragLeave(event)" ondrop="handleDrop(event, 'top')" onclick="handleBoxClick('top')">
                        <div class="camera-overlay-badge">
                            <span class="badge-dot" id="dotTop"></span>
                            <span id="labelTop">Top View (Arducam_TOP)</span>
                        </div>
                        <div class="camera-actions-overlay">
                            <button class="btn-pill-action" onclick="event.stopPropagation(); setViewMode('live_orig3')">Switch to Origin 3 ➔</button>
                            <button class="btn-pill-action" onclick="event.stopPropagation(); document.getElementById('fileTopInput').click()">📁 Upload Image</button>
                        </div>
                        <video id="videoTop" autoplay playsinline muted></video>
                        <img id="imgUploadTop" style="display: none;" alt="Top Preview" />
                        <canvas id="overlayTop" class="guide-canvas"></canvas>
                    </div>

                    <!-- Live Origin Camera Viewport (Full Screen) -->
                    <div class="camera-full-container" id="containerOrig" style="display: none;" ondragover="handleDragOver(event)" ondragleave="handleDragLeave(event)" ondrop="handleDrop(event, 'orig3')" onclick="handleBoxClick('orig3')">
                        <div class="camera-overlay-badge">
                            <span class="badge-dot" id="dotOrig"></span>
                            <span id="labelOrig">Origin View (Arducam_ORIGIN - Origin 3)</span>
                        </div>
                        <div class="camera-actions-overlay">
                            <button class="btn-pill-action" onclick="event.stopPropagation(); setViewMode('live_top')">⬅ Switch to Top View</button>
                            <button class="btn-pill-action" onclick="event.stopPropagation(); document.getElementById('fileOrig3Input').click()">📁 Upload Image</button>
                        </div>
                        <video id="videoOrig3" autoplay playsinline muted></video>
                        <img id="imgUploadOrig" style="display: none;" alt="Origin Preview" />
                        <canvas id="overlayOrig3" class="guide-canvas"></canvas>
                    </div>

                    <!-- Single Result View -->
                    <div class="single-view" id="singleViewArea">
                        <div class="camera-overlay-badge">
                            <span id="resultViewLabel">Inspection Result</span>
                        </div>
                        <div class="camera-actions-overlay">
                            <button class="btn-pill-action" onclick="replaceImageForCurrentTab()">📁 Replace Image</button>
                            <button class="btn-pill-action" onclick="setViewMode('live_top')">🔄 Back to Live</button>
                        </div>
                        <img id="singleResultImg" alt="Result Inspection View" />
                    </div>
                </div>
            </div>

            <!-- Bottom Minimal Controls -->
            <div class="bottom-bar">
                <div class="cam-controls">
                    <span>Top:</span>
                    <select id="selectCamTop" onchange="onCameraSelectionChanged()"></select>
                    <span>Origin:</span>
                    <select id="selectCamOrig" onchange="onCameraSelectionChanged()"></select>
                    <button class="btn-icon" onclick="swapCameras()" title="Swap Camera Inputs">⇄ Swap</button>
                    <button class="btn-icon" onclick="restartLiveCameras()" title="Refresh Camera Streams">🔄 Refresh</button>
                </div>
                <div id="footerStatus">Scan QR Code or enter Device Serial No. to validate with server and activate START.</div>
            </div>
        </div>

        <!-- RIGHT: Clean Minimal Sidebar -->
        <div class="sidebar-wrapper">
            <!-- Overall Verdict Banner -->
            <div class="verdict-box" id="verdictBanner">
                <div class="verdict-status" id="verdictText">AWAITING START</div>
                <div class="verdict-device" id="verdictDevice">Device: —</div>
                <div class="verdict-desc" id="verdictDesc">Scan QR code to validate with PMS server & activate START button</div>
            </div>

            <!-- Inspection Results Card -->
            <div class="sidebar-card" id="resultsSidebarCard">
                <div class="section-header">
                    <span>Inspection Summary</span>
                    <span style="font-size: 10px; color: var(--text-muted);">AI Validation</span>
                </div>

                <!-- Idle state: before inspection runs -->
                <div id="resultsIdleState" class="result-idle-state">
                    <span class="status-pill">—</span>
                    <span style="font-size: 12px; color: var(--text-muted); margin-left: 8px;">Awaiting Inspection</span>
                </div>

                <!-- Detailed Case Checklist Table (shows pass or fail for all cases) -->
                <div id="resultsGrid" style="display: none; margin-top: 10px;">
                    <div class="section-header">
                        <span>Wire Origins</span>
                        <span style="font-size: 10px; color: var(--text-muted);">Side & Top View</span>
                    </div>
                    <div class="check-table">
                        <div class="check-row">
                            <div class="check-name">Origin 3 (Grommet)</div>
                            <span class="status-pill" id="pillOrig3Grommet">—</span>
                        </div>
                        <div class="check-row">
                            <div class="check-name">Origin 1 (Bottom Left)</div>
                            <span class="status-pill" id="pillObl">—</span>
                        </div>
                        <div class="check-row">
                            <div class="check-name">Origin 2 (Upper Middle)</div>
                            <span class="status-pill" id="pillOum">—</span>
                        </div>
                        <div class="check-row">
                            <div class="check-name">Origin 3 (Bottom Right)</div>
                            <span class="status-pill" id="pillObr">—</span>
                        </div>
                    </div>

                    <div class="section-header" style="margin-top: 14px;">
                        <span>Wiring Routing</span>
                        <span style="font-size: 10px; color: var(--text-muted);">Top View</span>
                    </div>
                    <div class="check-table">
                        <div class="check-row">
                            <div class="check-name">Case 1: Bottom Left</div>
                            <span class="status-pill" id="pillC1">—</span>
                        </div>
                        <div class="check-row">
                            <div class="check-name">Case 2: Upper Middle</div>
                            <span class="status-pill" id="pillC2">—</span>
                        </div>
                        <div class="check-row">
                            <div class="check-name">Case 3: Bottom Right</div>
                            <span class="status-pill" id="pillC3">—</span>
                        </div>
                    </div>
                </div>

                <!-- Pass state: ONLY display Pass when the inspection passes -->
                <div id="resultsPassState" class="result-pass-state" style="display: none; margin-top: 12px;">
                    <div class="pass-badge-large">✔ PASS</div>
                </div>

                <!-- Fail state: In case of fail, tell what all things failed -->
                <div id="resultsFailState" class="result-fail-state" style="display: none; margin-top: 12px;">
                    <div class="fail-badge-large">✖ FAIL</div>
                    <div class="fail-list-title" style="margin-top: 6px;">Failed Checks:</div>
                    <div class="fail-items-container" id="failItemsContainer">
                        <!-- Populated dynamically with failed items -->
                    </div>
                </div>
            </div>

            <!-- Collapsible Diagnostic Log -->
            <details>
                <summary>Diagnostic Details</summary>
                <div class="diag-log" id="diagLog">Ready for inspection.</div>
            </details>
        </div>
    </div>

    <!-- Hidden File Inputs -->
    <input type="file" id="fileTopInput" accept="image/*" onchange="handleManualUpload(event, 'top')">
    <input type="file" id="fileOrig3Input" accept="image/*" onchange="handleManualUpload(event, 'orig3')">
    <input type="file" id="fileBothInput" accept="image/*" multiple onchange="handleBatchUpload(event)">

    <!-- JAVASCRIPT CONTROLLER -->
    <script>
        let streamTop = null;
        let streamOrig3 = null;
        let animFrameTop = null;
        let animFrameOrig3 = null;

        let rawTopBase64 = null;
        let rawOrig3Base64 = null;
        let uploadedTopBase64 = null;
        let uploadedOrigBase64 = null;

        let resTopWiresAnnotated = null;
        let resTopOriginsAnnotated = null;
        let resOrig3Annotated = null;

        let currentViewMode = 'live_top';
        let devicesList = [];

        const videoTop = document.getElementById('videoTop');
        const videoOrig3 = document.getElementById('videoOrig3');
        const imgUploadTop = document.getElementById('imgUploadTop');
        const imgUploadOrig = document.getElementById('imgUploadOrig');

        const overlayTop = document.getElementById('overlayTop');
        const overlayOrig3 = document.getElementById('overlayOrig3');
        const containerTop = document.getElementById('containerTop');
        const containerOrig = document.getElementById('containerOrig');
        const singleViewArea = document.getElementById('singleViewArea');
        const singleResultImg = document.getElementById('singleResultImg');

        const selectCamTop = document.getElementById('selectCamTop');
        const selectCamOrig = document.getElementById('selectCamOrig');
        const btnStartInspection = document.getElementById('btnStartInspection');
        const deviceSerialInput = document.getElementById('deviceSerialInput');

        
        function clearPreviousProductOutput() {
            // Clear the previous inspection images.
            resTopWiresAnnotated = null;
            resTopOriginsAnnotated = null;
            resOrig3Annotated = null;

            // Clear the single-result image, if displayed.
            if (singleResultImg) {
                singleResultImg.removeAttribute("src");
                singleResultImg.style.display = "none";
            }

            // Reset simplified inspection result card
            const resultsIdle = document.getElementById("resultsIdleState");
            const resultsGrid = document.getElementById("resultsGrid");
            const resultsPass = document.getElementById("resultsPassState");
            const resultsFail = document.getElementById("resultsFailState");
            const failItemsContainer = document.getElementById("failItemsContainer");
            if (resultsIdle) resultsIdle.style.display = "flex";
            if (resultsGrid) resultsGrid.style.display = "none";
            if (resultsPass) resultsPass.style.display = "none";
            if (resultsFail) resultsFail.style.display = "none";
            if (failItemsContainer) failItemsContainer.innerHTML = "";

            // Clear inspection result indicators (safe fallback)
            [
                "pillOrig3Grommet",
                "pillObl",
                "pillOum",
                "pillObr",
                "pillC1",
                "pillC2",
                "pillC3"
            ].forEach(id => {
                const element = document.getElementById(id);
                if (element) {
                    element.textContent = "--";
                    element.className = element.className
                        .replace(/\b(pass|fail|success|error)\b/gi, "")
                        .trim();
                }
            });

            // Hide the previous verdict banner.
            const verdictBanner = document.getElementById("verdictBanner");
            if (verdictBanner) {
                verdictBanner.style.display = "none";
            }

            // Clear the previous product's displayed identity and verdict.
            [
                "verdictText",
                "verdictDevice",
                "verdictDesc"
            ].forEach(id => {
                const element = document.getElementById(id);
                if (element) {
                    element.textContent = "";
                }
            });

            // Clear the previous product's uploaded input images.
            rawTopBase64 = null;
            rawOrig3Base64 = null;
            uploadedTopBase64 = null;
            uploadedOrigBase64 = null;

            // Clear the previous Top image preview.
            imgUploadTop.removeAttribute("src");
            imgUploadTop.style.display = "none";

            // Clear the previous Origin image preview.
            imgUploadOrig.removeAttribute("src");
            imgUploadOrig.style.display = "none";

            // Restore the camera elements.
            videoTop.style.display = "block";
            videoOrig3.style.display = "block";

            // Hide the empty result image.
            singleResultImg.removeAttribute("src");
            singleResultImg.style.display = "none";

            // Return to the normal camera/upload screen.
            setViewMode("live_top");

            // Keep the empty result image hidden.
            singleResultImg.removeAttribute("src");
            singleResultImg.style.display = "none";
        }


        function parseQrInput(text) {
            if (!text) return { id: '', sr: '' };
            const idMatch = text.match(/(?:^|[;,|\s])\s*(?:APoNumberId|APO(?:_?ID)?|Id)\s*[:=]\s*([^;,|]+)/i);
            const srMatch = text.match(/(?:^|[;,|\s])\s*(?:DeviceSerial(?:No)?|Serial(?:No)?|Sr(?:No)?|SN)\s*[:=]\s*([^;,|]+)/i);
            let id = idMatch ? idMatch[1].trim() : '';
            let sr = srMatch ? srMatch[1].trim() : '';
            if (!id && !sr) {
                for (const sep of [',', '|', ';', '/']) {
                    if (text.includes(sep)) {
                        const parts = text.split(sep).map(p => p.trim()).filter(Boolean);
                        if (parts.length >= 2) {
                            id = parts[0];
                            sr = parts[1];
                            break;
                        }
                    }
                }
                if (!sr && !id) {
                    sr = text.trim();
                }
            }
            return { id, sr };
        }

        let currentValidatedCode = null;
        let isValidating = false;

        
        async function validateScannedDevice() {
            const rawVal = deviceSerialInput.value.trim();

            if (!rawVal) {
                setValidationState('idle', 'Scan to Validate');
                btnStartInspection.disabled = true;
                btnStartInspection.classList.remove('ready-pulse');
                pmsPayloadsCache = null;
                document.getElementById('footerStatus').innerText =
                    "Scan QR Code or enter Device Serial No. to validate with server and activate START.";
                return;
            }

            // Prevent validation while another validation request is in progress.
            if (isValidating) return;

            // Check the application version before processing every new scan.
            try {
                const versionResponse = await fetch('/api/version', {
                    cache: 'no-store'
                });

                if (!versionResponse.ok) {
                    throw new Error("Could not retrieve the application version.");
                }

                const versionData = await versionResponse.json();
                const appVersion = String(versionData.version || '').trim();

                if (!versionData.success || !appVersion) {
                    currentValidatedCode = null;
                    pmsPayloadsCache = null;
                    btnStartInspection.disabled = true;
                    btnStartInspection.classList.remove('ready-pulse');

                    setValidationState('failed', 'Version Missing');

                    document.getElementById('footerStatus').innerText =
                        "⚠ Application version is missing. Scanning is blocked.";

                   alert("Application Version Missing! Please configure the application version in Admin Settings before starting inspection.");

                    return;
                }

            } catch (error) {
                currentValidatedCode = null;
                pmsPayloadsCache = null;
                btnStartInspection.disabled = true;
                btnStartInspection.classList.remove('ready-pulse');

                setValidationState('failed', 'Version Check Failed');

                document.getElementById('footerStatus').innerText =
                    "⚠ Could not verify the application version. Scanning is blocked.";

                alert(
                    "Unable to verify the application version, Please check the application connection and try scanning again");
                

                console.error("Application version check failed:", error);
                return;
            }

            // If already validated for this exact string and button is currently active.
            if (currentValidatedCode === rawVal && !btnStartInspection.disabled) {
                return;
            }

            isValidating = true;
            btnStartInspection.disabled = true;
            btnStartInspection.classList.remove('ready-pulse');
            pmsPayloadsCache = null;
            setValidationState('validating', '⏳ Sending Payloads...');

            document.getElementById('footerStatus').innerText =
                "Sending Whitelisting, SerialNo, and StatusValidation payloads to server. Waiting for CONFIG SUCCESSFUL...";

            try {
                const res = await fetch('/api/pms/validate', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({ device_id: rawVal })
                });

                const data = await res.json();

                if (res.ok && data.validated) {
                    currentValidatedCode = rawVal;

                    // Response received, UtcRequest_ID matched, CONFIG SUCCESSFUL.
                    pmsPayloadsCache = data.payloads;
                    pmsPayloadsCache.server_responses = data.server_responses || {};
                    pmsPayloadsCache._forDevice = rawVal;

                    setValidationState('success', '✔ CONFIG SUCCESSFUL');
                    btnStartInspection.disabled = false;
                    btnStartInspection.classList.add('ready-pulse');

                    document.getElementById('footerStatus').innerText =
                        `✔ Server responded CONFIG SUCCESSFUL for Serial [${data.serial_no}], APO [${data.apo_id}]. START button ACTIVATED.`;

                    showConfigResultPopup(true, {
                        title: "PMS Configuration Successful",
                        serial_no: data.serial_no,
                        apo_id: data.apo_id,
                        message: "All 3 validation payloads (Whitelisting, SerialNo, Status) confirmed CONFIG SUCCESSFUL by server. The START button is now ACTIVATED."
                    });

                } else {
                    currentValidatedCode = null;
                    pmsPayloadsCache = null;

                    setValidationState('failed', '✖ Validation Failed');
                    btnStartInspection.disabled = true;
                    btnStartInspection.classList.remove('ready-pulse');

                    const errMsg = data.message ||
                        'Server rejected configuration or request ID mismatched.';

                    document.getElementById('footerStatus').innerText =
                        `✖ PMS Validation Failed: ${errMsg}. START button remains disabled.`;

                    showConfigResultPopup(false, {
                        title: "PMS Configuration Failed",
                        serial_no: data.serial_no || rawVal,
                        apo_id: data.apo_id,
                        error: errMsg,
                        message: "Server rejected configuration or request ID mismatched. START button remains disabled and JSON fields stay empty."
                    });
                }

            } catch (err) {
                currentValidatedCode = null;
                pmsPayloadsCache = null;

                setValidationState('failed', '✖ Server Error');
                btnStartInspection.disabled = true;
                btnStartInspection.classList.remove('ready-pulse');

                document.getElementById('footerStatus').innerText =
                    `✖ Cannot reach PMS validation service: ${err.message}. START button disabled.`;

                showConfigResultPopup(false, {
                    title: "PMS Connection Error / Timeout",
                    serial_no: rawVal,
                    error: err.message,
                    message: "Cannot reach PMS validation service or server response timed out. START button remains disabled."
                });

            } finally {
                isValidating = false;
            }
        }


        function setValidationState(state, text) {
            const badge = document.getElementById('qrValidationBadge');
            if (!badge) return;
            badge.className = `qr-badge ${state}`;
            badge.innerText = text;
        }

        function focusQrInput() {
            if (deviceSerialInput) {
                deviceSerialInput.focus();
                if (deviceSerialInput.value) {
                    deviceSerialInput.select();
                }
            }
        }

        window.addEventListener('DOMContentLoaded', async () => {
            // Immediate auto-focus as soon as app opens or reloads
            focusQrInput();
            setTimeout(focusQrInput, 50);
            setTimeout(focusQrInput, 300);

            // Initial state: START button disabled until validation succeeds with CONFIG SUCCESSFUL
            btnStartInspection.disabled = true;
            btnStartInspection.classList.remove('ready-pulse');
            setValidationState('idle', 'Scan to Validate');
            pmsPayloadsCache = null;

            // QR scanner enters text into text box and fires Enter to validate
            deviceSerialInput.addEventListener('keydown', (e) => {
                if (e.key === 'Enter') {
                    e.preventDefault();
                    validateScannedDevice();
                }
            });

            // If input is changed or cleared after validation, disable START button immediately
            let inputDebounceTimer = null;
            deviceSerialInput.addEventListener('input', () => {
                const val = deviceSerialInput.value.trim();
                if (val !== currentValidatedCode) {
                    clearPreviousProductOutput();

                    btnStartInspection.disabled = true;
                    btnStartInspection.classList.remove('ready-pulse');
                    pmsPayloadsCache = null; // Keep JSON fields empty until server confirms
                    if (!val) {
                        setValidationState('idle', 'Scan to Validate');
                        document.getElementById('footerStatus').innerText = "Scan QR Code or enter Device Serial No. to validate with server and activate START.";
                    } else {
                        setValidationState('idle', 'Press Enter to Validate');
                        document.getElementById('footerStatus').innerText = "Device ID entered. Press Enter or scan to send payloads and validate with PMS server.";
                    }
                }

                // Auto-trigger validation if a completed QR string is scanned / pasted
                clearTimeout(inputDebounceTimer);
                if (val && (val.includes(';') || val.includes('|')) && (val.toLowerCase().includes('id:') || val.toLowerCase().includes('sr:'))) {
                    inputDebounceTimer = setTimeout(() => {
                        if (val !== currentValidatedCode) {
                            validateScannedDevice();
                        }
                    }, 350);
                }
            });

            await initCameraDevices();

            // Re-focus after camera stream initialization completes
            focusQrInput();
            setTimeout(focusQrInput, 500);

            // Keyboard shortcuts: Spacebar triggers inspection; Esc dismisses active toast
            window.addEventListener('keydown', (e) => {
                if (e.key === 'Escape') {
                    const activeToast = document.querySelector('.config-toast.show');
                    if (activeToast) {
                        dismissToast(activeToast);
                    }
                }
                if (e.code === 'Space' && e.target.tagName !== 'INPUT' && e.target.tagName !== 'TEXTAREA') {
                    e.preventDefault();
                    if (!btnStartInspection.disabled) {
                        triggerStartInspection();
                    } else {
                        document.getElementById('footerStatus').innerText = "⚠️ START is disabled: QR must be scanned and validated by server (CONFIG SUCCESSFUL) first.";
                    }
                }
            });
        });

        // Also ensure textbox is focused on window load and window focus/switch
        window.addEventListener('load', () => {
            focusQrInput();
            setTimeout(focusQrInput, 200);
        });

        window.addEventListener('focus', () => {
            focusQrInput();
        });

        // Auto re-focus textbox when clicking outside interactive controls
        document.addEventListener('click', (e) => {
            const tag = e.target.tagName;
            if (tag !== 'BUTTON' && tag !== 'SELECT' && tag !== 'INPUT' && tag !== 'TEXTAREA' && !e.target.closest('#pmsJsonModal') && !e.target.closest('#configToastContainer') && !e.target.closest('.modal-content')) {
                focusQrInput();
            }
        });

        function clearQrInput() {
            deviceSerialInput.value = '';
            currentValidatedCode = null;
            pmsPayloadsCache = null;
            btnStartInspection.disabled = true;
            btnStartInspection.classList.remove('ready-pulse');
            setValidationState('idle', 'Scan to Validate');
            document.getElementById('footerStatus').innerText = "Ready. Scan QR Code (Id: ..; Sr: ..;) to validate and activate START.";
            focusQrInput();
        }

        async function initCameraDevices() {
            try {
                const tempStream = await navigator.mediaDevices.getUserMedia({ video: true, audio: false });
                tempStream.getTracks().forEach(t => t.stop());

                const allDevices = await navigator.mediaDevices.enumerateDevices();
                devicesList = allDevices.filter(d => d.kind === 'videoinput');

                selectCamTop.innerHTML = '';
                selectCamOrig.innerHTML = '';

                const savedTopId = localStorage.getItem('selected_top_cam');
                const savedOrigId = localStorage.getItem('selected_orig_cam');

                let topIdx = 0;
                let origIdx = devicesList.length > 1 ? 1 : 0;

                devicesList.forEach((dev, idx) => {
                    const baseLabel = dev.label || `Camera ${idx + 1}`;
                    const label = devicesList.filter(d => (d.label || '') === dev.label).length > 1 
                        ? `${baseLabel} (#${idx + 1})` 
                        : baseLabel;
                    const lower = label.toLowerCase();

                    selectCamTop.add(new Option(label, dev.deviceId));
                    selectCamOrig.add(new Option(label, dev.deviceId));

                    if (savedTopId && dev.deviceId === savedTopId) {
                        topIdx = idx;
                    } else if (!savedTopId && (lower.includes('top') || lower.includes('arducam_top'))) {
                        topIdx = idx;
                    }

                    if (savedOrigId && dev.deviceId === savedOrigId) {
                        origIdx = idx;
                    } else if (!savedOrigId && (lower.includes('origin') || lower.includes('arducam_origin') || lower.includes('grommet'))) {
                        origIdx = idx;
                    }
                });

                if (devicesList.length > 0) {
                    selectCamTop.selectedIndex = topIdx;
                    selectCamOrig.selectedIndex = origIdx;
                    await startDualStreams();
                } else {
                    document.getElementById('footerStatus').innerText = "No cameras detected. Use Upload button to test images.";
                }
            } catch (err) {
                console.warn("Camera init error:", err);
                document.getElementById('footerStatus').innerText = "Camera access unavailable. Using upload mode.";
            }
        }

        async function startDualStreams() {
            stopAllStreams();

            const topDeviceId = selectCamTop.value;
            const origDeviceId = selectCamOrig.value;

            if (topDeviceId) localStorage.setItem('selected_top_cam', topDeviceId);
            if (origDeviceId) localStorage.setItem('selected_orig_cam', origDeviceId);

            if (!uploadedTopBase64) {
                try {
                    const constraintsTop = topDeviceId ? { video: { deviceId: { exact: topDeviceId }, width: 1280, height: 720 } } : { video: { width: 1280, height: 720 } };
                    streamTop = await navigator.mediaDevices.getUserMedia(constraintsTop);
                    videoTop.srcObject = streamTop;
                    videoTop.style.display = 'block';
                    imgUploadTop.style.display = 'none';
                    await videoTop.play();
                    document.getElementById('dotTop').classList.add('active');
                    renderTopGuides();
                } catch (e) {
                    console.warn("Top stream error:", e);
                    document.getElementById('dotTop').classList.remove('active');
                }
            }

            if (!uploadedOrigBase64) {
                try {
                    const constraintsOrig = origDeviceId ? { video: { deviceId: { exact: origDeviceId }, width: 1280, height: 720 } } : { video: { width: 1280, height: 720 } };
                    streamOrig3 = await navigator.mediaDevices.getUserMedia(constraintsOrig);
                    videoOrig3.srcObject = streamOrig3;
                    videoOrig3.style.display = 'block';
                    imgUploadOrig.style.display = 'none';
                    await videoOrig3.play();
                    document.getElementById('dotOrig').classList.add('active');
                    renderOrig3Guides();
                } catch (e) {
                    console.warn("Origin stream error:", e);
                    document.getElementById('dotOrig').classList.remove('active');
                }
            }

            setViewMode(currentViewMode || 'live_top');
        }

        function stopAllStreams() {
            if (streamTop) {
                streamTop.getTracks().forEach(t => t.stop());
                streamTop = null;
            }
            if (streamOrig3) {
                streamOrig3.getTracks().forEach(t => t.stop());
                streamOrig3 = null;
            }
            if (animFrameTop) cancelAnimationFrame(animFrameTop);
            if (animFrameOrig3) cancelAnimationFrame(animFrameOrig3);
        }

        function restartLiveCameras() {
            startDualStreams();
        }

        function onCameraSelectionChanged() {
            startDualStreams();
        }

        function swapCameras() {
            const temp = selectCamTop.selectedIndex;
            selectCamTop.selectedIndex = selectCamOrig.selectedIndex;
            selectCamOrig.selectedIndex = temp;
            startDualStreams();
        }

        function resetToLiveCameras() {
            uploadedTopBase64 = null;
            uploadedOrigBase64 = null;
            document.getElementById('labelTop').innerText = "Top View (Arducam_TOP)";
            document.getElementById('labelOrig').innerText = "Origin View (Arducam_ORIGIN - Origin 3)";
            setViewMode('live_top');
            startDualStreams();
            document.getElementById('footerStatus').innerText = "Restored Live Cameras.";
        }

        function renderTopGuides() {
            if (uploadedTopBase64) return;
            if (!streamTop || !videoTop.videoWidth) {
                animFrameTop = requestAnimationFrame(renderTopGuides);
                return;
            }
            overlayTop.width = videoTop.clientWidth;
            overlayTop.height = videoTop.clientHeight;
            const ctx = overlayTop.getContext('2d');
            ctx.clearRect(0, 0, overlayTop.width, overlayTop.height);

            const sx = overlayTop.width / 1024;
            const sy = overlayTop.height / 576;

            // Coordinate-based ROIs removed

            animFrameTop = requestAnimationFrame(renderTopGuides);
        }

        function renderOrig3Guides() {
            if (uploadedOrigBase64) return;
            if (!streamOrig3 || !videoOrig3.videoWidth) {
                animFrameOrig3 = requestAnimationFrame(renderOrig3Guides);
                return;
            }
            overlayOrig3.width = videoOrig3.clientWidth;
            overlayOrig3.height = videoOrig3.clientHeight;
            const ctx = overlayOrig3.getContext('2d');
            ctx.clearRect(0, 0, overlayOrig3.width, overlayOrig3.height);

            const sx = overlayOrig3.width / 1024;
            const sy = overlayOrig3.height / 576;

            // Coordinate-based ROIs removed

            animFrameOrig3 = requestAnimationFrame(renderOrig3Guides);
        }

        function drawGuideBox(ctx, sx1, sy1, sx2, sy2, color, text) {
            ctx.strokeStyle = color;
            ctx.lineWidth = 1.5;
            ctx.setLineDash([3, 3]);
            ctx.strokeRect(sx1, sy1, sx2 - sx1, sy2 - sy1);
            ctx.setLineDash([]);
            if (text) {
                ctx.font = '11px -apple-system, sans-serif';
                ctx.fillStyle = color;
                ctx.fillText(text, sx1 + 3, sy1 - 5);
            }
        }

        function grabFrameFromVideo(videoEl) {
            if (!videoEl || !videoEl.videoWidth) return null;
            const canvas = document.createElement('canvas');
            canvas.width = 1024;
            canvas.height = 576;
            const ctx = canvas.getContext('2d');
            ctx.drawImage(videoEl, 0, 0, 1024, 576);
            return canvas.toDataURL('image/jpeg', 0.95);
        }

        
        
        async function triggerStartInspection() {
            if (btnStartInspection.disabled) {
                return;
            }

            // Verify the application version before starting inspection.
            try {
                const versionResponse = await fetch('/api/version', {
                    cache: 'no-store'
                });

                if (!versionResponse.ok) {
                    throw new Error("Could not retrieve the application version.");
                }

                const versionData = await versionResponse.json();
                const appVersion = String(versionData.version || '').trim();

                if (!versionData.success || !appVersion) {
                    btnStartInspection.disabled = true;
                    btnStartInspection.classList.remove('ready-pulse');

                    document.getElementById('footerStatus').innerText =
                        "Application version is missing. Inspection cannot start.";

                    alert("Application Version Missing! Please configure the application version in Admin Settings before starting inspection.");

                    return;
                }

            } catch (error) {
                btnStartInspection.disabled = true;
                btnStartInspection.classList.remove('ready-pulse');

                document.getElementById('footerStatus').innerText =
                    "Could not verify the application version. Inspection is blocked.";

                alert('Unable to verify the application version, Please check the application connection and try again.');

                console.error("Application version check failed:", error);
                return;
            }

            // Existing inspection logic.
            btnStartInspection.disabled = true;
            btnStartInspection.classList.remove('ready-pulse');

            document.getElementById('footerStatus').innerText =
                "Inspecting both cameras and logging to Excel...";

            if (uploadedTopBase64) {
                rawTopBase64 = uploadedTopBase64;
            } else if (streamTop && videoTop.videoWidth) {
                rawTopBase64 = grabFrameFromVideo(videoTop);
            }

            if (uploadedOrigBase64) {
                rawOrig3Base64 = uploadedOrigBase64;
            } else if (streamOrig3 && videoOrig3.videoWidth) {
                rawOrig3Base64 = grabFrameFromVideo(videoOrig3);
            }

            if (!rawTopBase64 || !rawOrig3Base64) {
                alert(
                    "Please ensure both Top View and Origin View images are available via live cameras or upload."
                );

                btnStartInspection.disabled = false;
                return;
            }

            const currentContainer =
                currentViewMode === 'live_orig3' ? containerOrig : containerTop;

            currentContainer.classList.add('flash-anim');

            setTimeout(
                () => currentContainer.classList.remove('flash-anim'),
                300
            );

            const scannedDeviceId = deviceSerialInput.value.trim() ||
                `PCB-${new Date().toISOString().replace(/[-:T.]/g, '').slice(0, 14)}`;

            try {
                const res = await fetch('/api/inspect', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({
                        device_id: scannedDeviceId,
                        top_image: rawTopBase64,
                        orig3_image: rawOrig3Base64,
                        orig2_image: rawOrig3Base64
                    })
                });

                const data = await res.json();
                applyResults(data, scannedDeviceId);

            } catch (err) {
                console.error(err);

                document.getElementById('footerStatus').innerText =
                    "Inspection request failed.";

            } finally {
                btnStartInspection.disabled = true;
                btnStartInspection.classList.remove('ready-pulse');

                currentValidatedCode = null;

                setValidationState('idle', 'Scan Next QR');
                focusQrInput();
            }
        }


        
        
        function applyResults(data, deviceId) {
            resTopWiresAnnotated = data.top_wires_annotated;
            resTopOriginsAnnotated = data.top_origins_annotated;
            resOrig3Annotated = data.orig3_annotated || data.orig2_annotated;

            const r = data.results || {};

            // Final verdict comes from the backend after PMS validation.
            const finalVerdict = data.verdict || "SERVER UNREACHABLE";
            const pmsStatus = data.pms_report_status || "UNREACHABLE";
            const pmsResponse = data.pms_mqtt_response || {};

            // Safe update for any existing pills
            setPill('pillOrig3Grommet', r.orig3_grommet_pass ?? r.orig2_grommet_pass);
            setPill('pillObl', r.obl_pass);
            setPill('pillOum', r.oum_pass);
            setPill('pillObr', r.obr_pass);
            setPill('pillC1', r.c1_pass);
            setPill('pillC2', r.c2_pass);
            setPill('pillC3', r.c3_pass);

            const banner = document.getElementById('verdictBanner');
            const verdictText = document.getElementById('verdictText');
            const verdictDevice = document.getElementById('verdictDevice');
            const verdictDesc = document.getElementById('verdictDesc');

            verdictDevice.innerText = "";
            verdictDevice.style.display = "none";

            // Update Inspection Result Card
            const idleState = document.getElementById('resultsIdleState');
            const resultsGrid = document.getElementById('resultsGrid');
            const passState = document.getElementById('resultsPassState');
            const failState = document.getElementById('resultsFailState');
            const failContainer = document.getElementById('failItemsContainer');

            if (idleState) idleState.style.display = 'none';
            if (resultsGrid) resultsGrid.style.display = 'block';

            const isOverallPass = data.combined_pass === true && (finalVerdict === "PASS" || finalVerdict === "SERVER UNREACHABLE" || pmsStatus === "SKIPPED");

            if (isOverallPass) {
                // PASS CASE: ONLY display Pass
                if (passState) passState.style.display = 'flex';
                if (failState) failState.style.display = 'none';

                banner.className = "verdict-box pass";
                verdictText.innerText = "PASS";
                verdictDesc.innerText = "AI inspection passed and PMS accepted ReportStore.";
            } else {
                // FAIL CASE: Tell what all things failed
                if (passState) passState.style.display = 'none';
                if (failState) failState.style.display = 'flex';

                if (failContainer) {
                    failContainer.innerHTML = '';
                    let failures = [];

                    if (data.fail_details && Array.isArray(data.fail_details) && data.fail_details.length > 0) {
                        failures = data.fail_details;
                    } else if (data.fail_items && Array.isArray(data.fail_items) && data.fail_items.length > 0) {
                        failures = data.fail_items.map(t => {
                            const p = t.split(':');
                            return { item: p[0].trim(), reason: p.slice(1).join(':').trim() || "Check failed" };
                        });
                    } else {
                        if (data.is_top_pass === false || !r.c1_pass) {
                            failures.push({ item: "Top View (Wire Pose Model)", reason: "Red/black wire routing or keypoints missing" });
                        }
                        if (data.is_side_pass === false || !(r.orig3_grommet_pass ?? r.orig2_grommet_pass)) {
                            failures.push({ item: "Side View (Origin 3 Model)", reason: "Origin 3 grommet not detected" });
                        }
                        if (pmsStatus === "REJECTED") {
                            failures.push({ item: "PMS ReportStore", reason: data.pms_rejection_reason || pmsResponse.reason || "Rejected by PMS" });
                        } else if (finalVerdict === "SERVER UNREACHABLE" || pmsStatus === "UNREACHABLE") {
                            failures.push({ item: "PMS Connection", reason: "Server unreachable / no MQTT response" });
                        }
                    }

                    if (failures.length === 0) {
                        failures.push({ item: "Inspection", reason: "AI validation criteria not met" });
                    }

                    failures.forEach(f => {
                        const itemDiv = document.createElement('div');
                        itemDiv.className = 'fail-item-row';
                        itemDiv.innerHTML = `
                            <span class="fail-icon">✖</span>
                            <div class="fail-text">
                                <div class="fail-name">${f.item || f.name || "Failed Check"}</div>
                                <div class="fail-subtext">${f.reason || "Check failed"}</div>
                            </div>
                        `;
                        failContainer.appendChild(itemDiv);
                    });
                }

                banner.className = "verdict-box fail";
                if (finalVerdict === "SERVER UNREACHABLE") {
                    verdictText.innerText = "SERVER UNREACHABLE";
                    verdictDesc.innerText =
                        pmsResponse.message ||
                        "No matching response received from PMS. Please check the server and MQTT connection.";

                } else if (pmsStatus === "REJECTED") {
                    verdictText.innerText = "FAIL";
                    const reason =
                        data.pms_rejection_reason ||
                        pmsResponse.reason ||
                        pmsResponse.request_status ||
                        "PMS rejected the ReportStore request.";
                    verdictDesc.innerText = `PMS rejected ReportStore: ${reason}`;

                } else {
                    verdictText.innerText = "FAIL";
                    verdictDesc.innerText = "AI inspection failed. PMS accepted ReportStore.";
                }
            }

            banner.style.display = "flex";

            console.log("Final verdict:", data.verdict);
            console.log("PMS status:", data.pms_report_status);
            console.log("Verdict banner:", banner);
            console.log("Verdict text:", verdictText.innerText);
            console.log("Banner display:", banner.style.display);

            if (data.notes) {
                document.getElementById('diagLog').innerText = data.notes;
            }

            // Cache the ReportStore payload generated for this run.
            if (data.pms_report_json) {
                if (!pmsPayloadsCache) {
                    pmsPayloadsCache = {};
                }

                pmsPayloadsCache.ReportStore = data.pms_report_json;
            }

            document.getElementById('footerStatus').innerText =
                `Inspection finished (${verdictText.innerText}). ` +
                `Device [${deviceId}]. PMS status: ${pmsStatus}.`;

            setViewMode('res_top_wires');
            // Make sure the previous image-clearing logic does not keep
            // the result image hidden.
            if (singleResultImg && singleResultImg.getAttribute('src')) {
                singleResultImg.style.display = "block";
            }
        }


        async function openExcelSheet() {
            try {
                const res = await fetch('/api/open_excel', { method: 'POST' });
                const data = await res.json();
                if (data.status === 'success') {
                    document.getElementById('footerStatus').innerText = "Opened Result/Inspection_Log.xlsx.";
                    window.location.href = '/api/download_excel';
                } else {
                    alert(data.message || "Could not open Excel sheet.");
                }
            } catch (e) {
                window.location.href = '/api/download_excel';
            }
        }

        function setPill(id, isPass) {
            const el = document.getElementById(id);
            if (!el) return;
            if (isPass) {
                el.innerText = "PASS";
                el.className = "status-pill pass";
            } else {
                el.innerText = "FAIL";
                el.className = "status-pill fail";
            }
        }

        function setViewMode(mode) {
            currentViewMode = mode;
            document.querySelectorAll('.nav-pill').forEach(btn => btn.classList.remove('active'));

            if (mode === 'live_top') {
                document.getElementById('pillLiveTop').classList.add('active');
                containerTop.style.display = 'flex';
                containerOrig.style.display = 'none';
                singleViewArea.style.display = 'none';
            } else if (mode === 'live_orig3' || mode === 'live_orig2') {
                document.getElementById('pillLiveOrig').classList.add('active');
                containerTop.style.display = 'none';
                containerOrig.style.display = 'flex';
                singleViewArea.style.display = 'none';
            } else {
                containerTop.style.display = 'none';
                containerOrig.style.display = 'none';
                singleViewArea.style.display = 'flex';

                if (mode === 'res_top_wires') {
                    document.getElementById('pillTopWires').classList.add('active');
                    document.getElementById('resultViewLabel').innerText = "Result: Top View Wire Routing Paths";
                    singleResultImg.src = resTopWiresAnnotated || rawTopBase64 || '';
                } else if (mode === 'res_top_origins') {
                    document.getElementById('pillTopOrigins').classList.add('active');
                    document.getElementById('resultViewLabel').innerText = "Result: Top View Wire Origins";
                    singleResultImg.src = resTopOriginsAnnotated || rawTopBase64 || '';
                } else if (mode === 'res_orig3' || mode === 'res_orig2') {
                    document.getElementById('pillOrig3').classList.add('active');
                    document.getElementById('resultViewLabel').innerText = "Result: Origin 3 Dedicated Grommet";
                    singleResultImg.src = resOrig3Annotated || rawOrig3Base64 || '';
                }
            }
        }

        function replaceImageForCurrentTab() {
            if (currentViewMode === 'res_orig3' || currentViewMode === 'res_orig2') {
                document.getElementById('fileOrig3Input').click();
            } else {
                document.getElementById('fileTopInput').click();
            }
        }

        function handleBoxClick(target) {
            if (target === 'top') {
                document.getElementById('fileTopInput').click();
            } else {
                document.getElementById('fileOrig3Input').click();
            }
        }

        function handleDragOver(e) {
            e.preventDefault();
            e.currentTarget.classList.add('drag-over');
        }

        function handleDragLeave(e) {
            e.preventDefault();
            e.currentTarget.classList.remove('drag-over');
        }

        function handleDrop(e, target) {
            e.preventDefault();
            e.currentTarget.classList.remove('drag-over');
            const files = e.dataTransfer.files;
            if (!files || files.length === 0) return;

            if (files.length >= 2) {
                loadFile(files[0], 'top');
                loadFile(files[1], 'orig3');
            } else {
                loadFile(files[0], target);
            }
        }

        function handleBatchUpload(e) {
            const files = e.target.files;
            if (!files || files.length === 0) return;

            if (files.length === 1) {
                loadFile(files[0], (currentViewMode === 'live_orig3' || currentViewMode === 'live_orig2') ? 'orig3' : 'top');
            } else {
                let topFile = files[0];
                let origFile = files[1];
                const n0 = files[0].name.toLowerCase();
                if (n0.includes('orig') || n0.includes('grommet') || n0.includes('front')) {
                    origFile = files[0];
                    topFile = files[1];
                }
                loadFile(topFile, 'top');
                loadFile(origFile, 'orig3');
            }
        }

        function handleManualUpload(e, target) {
            const file = e.target.files[0];
            if (!file) return;
            loadFile(file, target);
        }

        function loadFile(file, target) {
            const reader = new FileReader();
            reader.onload = function(evt) {
                const data = evt.target.result;
                if (target === 'top') {
                    uploadedTopBase64 = data;
                    rawTopBase64 = data;
                    videoTop.style.display = 'none';
                    imgUploadTop.src = data;
                    imgUploadTop.style.display = 'block';
                    document.getElementById('labelTop').innerText = `Top: ${file.name.slice(0, 14)}`;
                    document.getElementById('dotTop').classList.add('active');
                } else {
                    uploadedOrigBase64 = data;
                    rawOrig3Base64 = data;
                    videoOrig3.style.display = 'none';
                    imgUploadOrig.src = data;
                    imgUploadOrig.style.display = 'block';
                    document.getElementById('labelOrig').innerText = `Origin 3: ${file.name.slice(0, 14)}`;
                    document.getElementById('dotOrig').classList.add('active');
                }

                if (currentViewMode !== 'live_top' && currentViewMode !== 'live_orig3') {
                    setViewMode(currentViewMode);
                }
                document.getElementById('footerStatus').innerText = `Loaded ${file.name}. Ready to inspect.`;
            };
            reader.readAsDataURL(file);
        }

        // PMS JSON Modal Logic
        let pmsPayloadsCache = null;
        let activePmsTab = 'report';

        async function openPmsJsonModal() {
            const modal = document.getElementById('pmsJsonModal');
            if (modal) modal.style.display = 'flex';
            const rawVal = document.getElementById('deviceSerialInput') ? document.getElementById('deviceSerialInput').value.trim() : '';

            try {
                if (!pmsPayloadsCache || pmsPayloadsCache._forDevice !== rawVal) {
                    const res = await fetch(`/api/pms/payloads?device_id=${encodeURIComponent(rawVal)}`);
                    const data = await res.json();
                    const existingReport = pmsPayloadsCache ? pmsPayloadsCache.ReportStore : null;
                    if (data && data.payloads) {
                        pmsPayloadsCache = data.payloads;
                        pmsPayloadsCache.server_responses = data.server_responses || {};
                        pmsPayloadsCache._forDevice = rawVal;
                        // If we have an actual completed inspection report cached (not empty), retain it
                        if (existingReport && existingReport.ColumnsName && existingReport.ColumnsName[0] && existingReport.ColumnsName[0].TaktTime !== "") {
                            pmsPayloadsCache.ReportStore = existingReport;
                        }
                    }
                }
            } catch (err) {
                console.error("Failed to load PMS JSON:", err);
                const pre = document.getElementById('pmsJsonPre');
                if (pre) pre.innerText = "Failed to load PMS JSON: " + err.message;
            }
            renderPmsModalJson();
        }

        function closePmsJsonModal(e) {
            if (!e || e.target.id === 'pmsJsonModal' || (e.target.classList && (e.target.classList.contains('modal-close') || e.target.classList.contains('btn-modal-close')))) {
                const modal = document.getElementById('pmsJsonModal');
                if (modal) modal.style.display = 'none';
            }
        }

        function switchModalTab(tab) {
            activePmsTab = tab;
            document.querySelectorAll('.modal-tab').forEach(b => b.classList.remove('active'));
            if (tab === 'report') document.getElementById('tabReport')?.classList.add('active');
            if (tab === 'whitelist') document.getElementById('tabWhitelist')?.classList.add('active');
            if (tab === 'serial') document.getElementById('tabSerial')?.classList.add('active');
            if (tab === 'status') document.getElementById('tabStatus')?.classList.add('active');
            if (tab === 'responses') document.getElementById('tabResponses')?.classList.add('active');
            renderPmsModalJson();
        }

        function renderPmsModalJson() {
            const pre = document.getElementById('pmsJsonPre');
            if (!pre) return;
            if (!pmsPayloadsCache) {
                pre.innerText = "Loading PMS JSON payloads...";
                return;
            }
            let jsonTarget = pmsPayloadsCache.ReportStore;
            if (activePmsTab === 'whitelist') jsonTarget = pmsPayloadsCache.DeviceAutoWhitelisting;
            if (activePmsTab === 'serial') jsonTarget = pmsPayloadsCache.SerialNoValidation;
            if (activePmsTab === 'status') jsonTarget = pmsPayloadsCache.StatusValidation;
            if (activePmsTab === 'responses') jsonTarget = pmsPayloadsCache.server_responses || {};
            pre.innerText = JSON.stringify(jsonTarget || {}, null, 2);
        }

        function copyPmsJson() {
            const txt = document.getElementById('pmsJsonPre').innerText;
            navigator.clipboard.writeText(txt).then(() => {
                const btn = document.querySelector('.btn-copy-json');
                const orig = btn.innerText;
                btn.innerText = '✔ Copied!';
                setTimeout(() => btn.innerText = orig, 1500);
            });
        }

        // Slide In/Out Config Toast Popup for Success or Failure
        function showConfigResultPopup(isSuccess, info) {
            const container = document.getElementById('configToastContainer');
            if (!container) return;

            // Remove previous toasts to keep display clean and focused
            container.innerHTML = '';

            const toast = document.createElement('div');
            toast.className = `config-toast ${isSuccess ? 'success' : 'failed'}`;

            const duration = isSuccess ? 3200 : 4200; // 3.2s on success, 4.2s on error

            const title = info.title || (isSuccess ? 'PMS Configuration Successful' : 'PMS Configuration Failed');
            const badgeText = isSuccess ? 'CONFIG SUCCESSFUL' : 'CONFIG FAILED';
            const bodyMsg = isSuccess
                ? (info.message || 'All 3 validation payloads confirmed. START button ACTIVATED.')
                : (info.error ? `Reason: ${info.error}` : (info.message || 'Server rejected configuration.'));

            const serialText = info.serial_no ? `<span>Serial: <strong>${escapeHtml(info.serial_no)}</strong></span>` : '';
            const apoText = info.apo_id ? `<span>APO: <strong>${escapeHtml(info.apo_id)}</strong></span>` : '';
            const metaHtml = (serialText || apoText) ? `<div class="toast-meta">${serialText}${apoText}</div>` : '';

            toast.innerHTML = `
                <div class="toast-content">
                    <div class="toast-icon ${isSuccess ? 'success' : 'failed'}">
                        ${isSuccess ? '✔' : '✖'}
                    </div>
                    <div class="toast-text">
                        <div class="toast-header-row">
                            <span class="toast-title">${escapeHtml(title)}</span>
                            <span class="toast-badge ${isSuccess ? 'success' : 'failed'}">${badgeText}</span>
                            <button class="toast-close-btn" onclick="dismissToast(this.closest('.config-toast'))" title="Dismiss">&times;</button>
                        </div>
                        <div class="toast-body">${escapeHtml(bodyMsg)}</div>
                        ${metaHtml}
                    </div>
                </div>
                <div class="toast-progress">
                    <div class="toast-progress-bar ${isSuccess ? 'success' : 'failed'}" style="animation: toastProgressCountdown ${duration / 1000}s linear forwards;"></div>
                </div>
            `;

            container.appendChild(toast);

            // Slide in smoothly
            requestAnimationFrame(() => {
                toast.classList.add('show');
            });

            // Automatically slide out and remove after short amount of time
            const hideTimer = setTimeout(() => {
                dismissToast(toast);
            }, duration);
            toast._hideTimer = hideTimer;

            // Click to dismiss early if desired
            toast.addEventListener('click', (e) => {
                if (e.target.tagName !== 'BUTTON') {
                    dismissToast(toast);
                }
            });

            // Textbox always remains active and focused
            focusQrInput();
        }

        function dismissToast(toast) {
            if (!toast || toast.classList.contains('hide')) return;
            clearTimeout(toast._hideTimer);
            toast.classList.remove('show');
            toast.classList.add('hide');
            setTimeout(() => {
                if (toast && toast.parentNode) {
                    toast.parentNode.removeChild(toast);
                }
            }, 360);
            focusQrInput();
        }

        function escapeHtml(str) {
            return String(str || '').replace(/[&<>"']/g, (m) => ({
                '&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#39;'
            }[m]));
        }

        // MQTT connection status indicator
        
        // MQTT connection status indicator
        async function updateMqttStatus() {
            const indicator = document.getElementById('mqttStatus');
            const statusText = document.getElementById('mqttStatusText');

            if (!indicator || !statusText) return;

            try {
                const response = await fetch('/api/pms/status', {
                    cache: 'no-store'
                });

                if (!response.ok) {
                    throw new Error('Status request failed');
                }

                const data = await response.json();
                const isConnected =
                    data.connected === true &&
                    data.subscribed === true;

                indicator.classList.toggle('connected', isConnected);
                indicator.classList.toggle('disconnected', !isConnected);

                statusText.textContent = isConnected
                    ? 'SERVER Connected'
                    : 'SERVER Disconnected';

            } catch (error) {
                console.error('MQTT status update failed:', error);

                const currentIndicator =
                    document.getElementById('mqttStatus');
                const currentStatusText =
                    document.getElementById('mqttStatusText');

                if (currentIndicator && currentStatusText) {
                    currentIndicator.classList.remove('connected');
                    currentIndicator.classList.add('disconnected');
                    currentStatusText.textContent = 'SERVER Disconnected';
                }
            }
        }

        // Make the updater accessible and refresh every 2 seconds.
        window.updateMqttStatus = updateMqttStatus;

        window.updateMqttStatus();
        window.setInterval(window.updateMqttStatus, 2000);



        // ================= ADMIN VERSION SETTINGS =================

        
        function openVersionAdminModal() {
            const modal = document.getElementById("versionAdminModal");
            const message = document.getElementById("versionAdminMessage");

            // Reset the form whenever the modal opens.
            document.getElementById("adminAction").value = "version";
            document.getElementById("newAppVersion").value = "";
            document.getElementById("newJigUid").value = "";
            document.getElementById("newProductionLineNo").value = "";
            document.getElementById("adminVersionPassword").value = "";

            message.style.display = "none";
            message.textContent = "";

            // Default to the application version action.
            handleAdminActionChange();

            modal.style.display = "flex";
        }


        function closeVersionAdminModal(event) {
            if (event && event.target !== event.currentTarget) {
                return;
            }

            document.getElementById("versionAdminModal").style.display = "none";
        }

        
        
        
        async function submitAdminSettings() {
            const action = document.getElementById("adminAction").value;
            const password = document.getElementById("adminVersionPassword").value;
            const message = document.getElementById("versionAdminMessage");

            let endpoint = "";
            let payload = {};

            if (!password.trim()) {
                message.style.display = "block";
                message.style.color = "#dc2626";
                message.textContent = "Please enter the admin password.";
                return;
            }

            if (action === "version") {
                const version = document.getElementById("newAppVersion").value.trim();

                if (!version) {
                    message.style.display = "block";
                    message.style.color = "#dc2626";
                    message.textContent = "Please enter an application version.";
                    return;
                }

                endpoint = "/api/admin/version";
                payload = {
                    version: version,
                    password: password
                };

                message.textContent = "Updating application version...";

            } else if (action === "jig_uid") {
                const jigUid = document.getElementById("newJigUid").value.trim();

                if (!jigUid) {
                    message.style.display = "block";
                    message.style.color = "#dc2626";
                    message.textContent = "Please enter a Jig UID.";
                    return;
                }

                endpoint = "/api/admin/jig-uid";
                payload = {
                    jig_uid: jigUid,
                    password: password
                };

                message.textContent = "Updating Jig UID...";

            } else if (action === "production_line") {
                const productionLineNo = document
                    .getElementById("newProductionLineNo")
                    .value.trim();

                if (!productionLineNo) {
                    message.style.display = "block";
                    message.style.color = "#dc2626";
                    message.textContent =
                        "Please enter a GTI Version .";
                    return;
                }

                endpoint = "/api/admin/production-line";
                payload = {
                    production_line_no: productionLineNo,
                    password: password
                };

                message.textContent = "Updating GTI Version Number...";

            } else {
                return;
            }

            message.style.display = "block";
            message.style.color = "#d97706";

            try {
                const response = await fetch(endpoint, {
                    method: "POST",
                    headers: {
                        "Content-Type": "application/json"
                    },
                    body: JSON.stringify(payload)
                });

                const result = await response.json();

                if (!response.ok || !result.success) {
                    message.style.color = "#dc2626";
                    message.textContent =
                        result.message || "Failed to save settings.";
                    return;
                }

                if (action === "version") {
                    // Update only the application version badge.
                    const headerVersion =
                        document.getElementById("appVersionBadge");

                    if (headerVersion) {
                        const appVersion = String(result.version || "").trim();

                        headerVersion.textContent = appVersion
                            ? "Version: " + appVersion
                            : "";
                    }

                    // Do not modify startVersionBadge here.
                    // It is reserved for the Production Line Number.

                    message.textContent =
                        result.message ||
                        "Application version updated successfully.";

                } else if (action === "jig_uid") {
                    message.textContent =
                        result.message || "Jig UID updated successfully.";

                    // Invalidate cached payloads so the next request can
                    // fetch payloads using the updated configuration.
                    pmsPayloadsCache = null;

                } else if (action === "production_line") {
                    message.textContent =
                        result.message ||
                        "GTI Version Number updated successfully.";

                    // Refresh the displayed production line number.
                    if (typeof loadCurrentVersion === "function") {
                        await loadCurrentVersion();
                    }
                }

                message.style.color = "#16a34a";
                document.getElementById("adminVersionPassword").value = "";

            } catch (error) {
                console.error("Admin settings error:", error);
                message.style.color = "#dc2626";
                message.textContent =
                    "Could not connect to the application server.";
            }
        }




        
        function handleAdminActionChange() {
            const action = document.getElementById("adminAction").value;

            const versionGroup = document.getElementById("versionFieldGroup");
            const jigUidGroup = document.getElementById("jigUidFieldGroup");
            const productionLineGroup = document.getElementById("productionLineFieldGroup");
            const message = document.getElementById("versionAdminMessage");

            // Show only the field for the selected action.
            versionGroup.style.display =
                action === "version" ? "block" : "none";

            jigUidGroup.style.display =
                action === "jig_uid" ? "block" : "none";

            productionLineGroup.style.display =
                action === "production_line" ? "block" : "none";

            // Clear previous messages when changing actions.
            message.style.display = "none";
            message.textContent = "";

            // Clear the password so it is entered for the selected action.
            document.getElementById("adminVersionPassword").value = "";
        }


        
        
        async function loadCurrentVersion() {
            try {
                // Load the Application Version.
                const response = await fetch("/api/version", {
                    cache: "no-store"
                });

                if (!response.ok) {
                    throw new Error("Failed to fetch application version.");
                }

                const result = await response.json();

                // Update the Application Version badge.
                // If no version is configured, keep the badge blank.
                const versionBadge =
                    document.getElementById("appVersionBadge");

                if (versionBadge) {
                    const appVersion = String(result.version || "").trim();

                    versionBadge.textContent =
                        result.success && appVersion
                            ? "Version: " + appVersion
                            : "";
                }

                // Load the Production Line Number.
                const lineResponse = await fetch("/api/production-line", {
                    cache: "no-store"
                });

                if (!lineResponse.ok) {
                    throw new Error("Failed to fetch GTI Version.");
                }

                const lineResult = await lineResponse.json();

                // Update the Production Line Number badge.
                const productionLineBadge =
                    document.getElementById("startVersionBadge");

                if (productionLineBadge) {
                    productionLineBadge.textContent =
                        lineResult.success &&
                        lineResult.production_line_no
                            ? "GTI Version: " + lineResult.production_line_no
                            : "";
                }

            } catch (error) {
                console.error(
                    "Could not load dashboard settings:",
                    error
                );
            }
        }

        // Load the current version and Production Line Number when the page opens.
        loadCurrentVersion();


    </script>

    <!-- PMS JSON Modal -->
    <div id="pmsJsonModal" class="modal-backdrop" style="display:none;" onclick="closePmsJsonModal(event)">
        <div class="modal-box" onclick="event.stopPropagation()">
            <div class="modal-header">
                <h3>📡 Selec Controls - PMS JSON Output</h3>
                <button class="modal-close" onclick="closePmsJsonModal()">&times;</button>
            </div>
            <div class="modal-tabs">
                <button class="modal-tab active" id="tabReport" onclick="switchModalTab('report')">4. ReportStore (Test Result)</button>
                <button class="modal-tab" id="tabWhitelist" onclick="switchModalTab('whitelist')">1. Whitelisting</button>
                <button class="modal-tab" id="tabSerial" onclick="switchModalTab('serial')">2. Serial Validation</button>
                <button class="modal-tab" id="tabStatus" onclick="switchModalTab('status')">3. Status Validation</button>
                <button class="modal-tab" id="tabResponses" onclick="switchModalTab('responses')">5. Server Responses</button>
            </div>
            <pre id="pmsJsonPre" class="modal-pre">Loading JSON...</pre>
            <div class="modal-footer">
                <button class="btn-copy-json" onclick="copyPmsJson()">📋 Copy JSON</button>
                <button class="btn-modal-close" onclick="closePmsJsonModal()">Close</button>
                
            </div>
        </div>
    </div>
    

    <!-- Admin Settings Modal -->
    <div id="versionAdminModal"
        class="modal-backdrop"
        style="display: none;"
        onclick="closeVersionAdminModal(event)">

        <div class="modal-box" onclick="event.stopPropagation()">

            <div class="modal-header">
                <h3>⚙️ Admin Settings</h3>
                <button class="modal-close"
                        type="button"
                        onclick="closeVersionAdminModal()">
                    &times;
                </button>
            </div>
    
            <div style="padding: 20px;">

                <!-- Action Selector -->
                <label for="adminAction"
                    style="display: block; margin-bottom: 8px;">
                    Select Action
                </label>

                <select id="adminAction"
                        onchange="handleAdminActionChange()"
                        style="width: 100%; box-sizing: border-box;
                            padding: 10px; margin-bottom: 18px;
                            background: #1a2333; color: #f1f5f9;
                            border: 1px solid #334155; border-radius: 6px;">

                    <option value="version">Update Product Version</option>
                    <option value="jig_uid">Update Jig UID</option>
                    <option value="production_line">Update GTI Version</option>

                </select>

                <!-- Application Version Field -->
                <div id="versionFieldGroup">

                    <label for="newAppVersion"
                        style="display: block; margin-bottom: 8px;">
                        New Product Version
                    </label>

                    <input type="text"
                        id="newAppVersion"
                        placeholder="Example: v2.0"
                        maxlength="30"
                        autocomplete="off"
                        style="width: 100%; box-sizing: border-box;
                                padding: 10px; margin-bottom: 16px;">

                </div>

                <!-- Jig UID Field -->
                <div id="jigUidFieldGroup" style="display: none;">

                    <label for="newJigUid"
                        style="display: block; margin-bottom: 8px;">
                        New Jig UID
                    </label>

                    <input type="text"
                        id="newJigUid"
                        placeholder="Enter the new Jig UID"
                        maxlength="100"
                        autocomplete="off"
                        style="width: 100%; box-sizing: border-box;
                                padding: 10px; margin-bottom: 16px;">

                </div>

                <div id="productionLineFieldGroup" style="display: none;">
                    <label for="newProductionLineNo" style="display: block; margin-bottom: 8px;">
                        New GTI Version
                    </label>
                    <input type="text" id="newProductionLineNo"
                        placeholder="Enter the GTI Version"
                        maxlength="30" autocomplete="off"
                        style="width: 100%; box-sizing: border-box; padding: 10px; margin-bottom: 16px;">
                </div>

                <!-- Admin Password -->
                <label for="adminVersionPassword"
                    style="display: block; margin-bottom: 8px;">
                    Admin Password
                </label>

                <input type="password"
                    id="adminVersionPassword"
                    placeholder="Enter admin password"
                    autocomplete="current-password"
                    style="width: 100%; box-sizing: border-box;
                            padding: 10px;">

                <!-- Success / Error Message -->
                <p id="versionAdminMessage"
                role="status"
                aria-live="polite"
                style="display: none; margin-top: 12px;">
                </p>

            </div>

            <div class="modal-footer">

                <button type="button"
                        class="btn-modal-close"
                        onclick="closeVersionAdminModal()">
                    Cancel
                </button>

                <button type="button"
                        class="btn-start"
                        onclick="submitAdminSettings()">
                    Save Settings
                </button>

            </div>

        </div>
    </div>



    <!-- Slide In/Out Config Toast Container -->
    <div id="configToastContainer"></div>

    <!-- pms_config.json ERROR POPUP -->
    <div id="configErrorOverlay" style="display:none; position:fixed; inset:0; z-index:100000; background:rgba(0,0,0,0.78); align-items:center; justify-content:center;">
        <div style="background:var(--bg-surface-elevated); border:2px solid #ef4444; border-radius:14px; width:92%; max-width:680px; max-height:85vh; overflow:auto; padding:24px; color:var(--text-main); box-shadow:0 20px 60px rgba(0,0,0,0.6);">
            <h2 style="margin:0 0 8px 0; color:#f87171; font-size:20px;">&#9888; pms_config.json has an error</h2>
            <p style="margin:0 0 14px 0; color:var(--text-secondary); font-size:14px;">
                Fix the file and save it. This window closes automatically once the file is valid.
            </p>
            <pre id="configErrorText" style="white-space:pre-wrap; word-break:break-word; background:var(--bg-base); border:1px solid #7f1d1d; border-radius:8px; padding:14px; margin:0 0 16px 0; font-family:Consolas, Menlo, monospace; font-size:13px; color:#fecaca;"></pre>
            <div style="display:flex; gap:10px; justify-content:flex-end;">
                <button type="button" id="configErrorRecheck" style="padding:9px 18px; border-radius:8px; border:none; background:#2563eb; color:#fff; font-weight:600; cursor:pointer;">Re-check now</button>
                <button type="button" id="configErrorDismiss" style="padding:9px 18px; border-radius:8px; border:1px solid var(--text-muted); background:transparent; color:var(--text-main); cursor:pointer;">Dismiss</button>
            </div>
        </div>
    </div>

    <script>
        // Watches pms_config.json: pops up the exact error if it is invalid,
        // on page load and every few seconds while the app is running.
        (function () {
            var overlay = document.getElementById('configErrorOverlay');
            var textBox = document.getElementById('configErrorText');
            var currentError = null;
            var dismissedError = null;

            function showError(message) {
                currentError = message;
                textBox.textContent = message;
                overlay.style.display = 'flex';
            }

            function hideError() {
                overlay.style.display = 'none';
                currentError = null;
                dismissedError = null;
            }

            async function checkConfig(manual) {
                try {
                    var response = await fetch('/api/config/validate', { cache: 'no-store' });
                    var data = await response.json();

                    if (data.valid) {
                        hideError();
                    } else {
                        var message = data.error || 'Unknown error in pms_config.json.';
                        // Re-show if manual check, or the error is new/different.
                        if (manual || message !== dismissedError) {
                            showError(message);
                        }
                    }
                } catch (err) {
                    // Server temporarily unreachable - ignore, try again next tick.
                }
            }

            document.getElementById('configErrorRecheck').addEventListener('click', function () {
                checkConfig(true);
            });

            document.getElementById('configErrorDismiss').addEventListener('click', function () {
                dismissedError = currentError;
                overlay.style.display = 'none';
            });

            checkConfig(false);
            window.setInterval(function () { checkConfig(false); }, 3000);
        })();
    </script>
</body>
</html>
"""

@app.route("/")
def index():
    return render_template_string(HTML_PAGE)

@app.route("/api/pms/status", methods=["GET"])
def pms_status_endpoint():
    return jsonify(pms_mqtt.get_connection_status())

@app.route("/api/open_excel", methods=["POST"])
def open_excel_endpoint():
    try:
        if os.path.exists(EXCEL_PATH):
            os.startfile(EXCEL_PATH)
            return jsonify({"status": "success", "message": "Excel file opened."})
        elif os.path.exists(CSV_PATH):
            os.startfile(CSV_PATH)
            return jsonify({"status": "success", "message": "CSV log opened."})
        else:
            return jsonify({"status": "error", "message": "No inspection log generated yet. Run an inspection first."})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)})

@app.route("/api/download_excel")
def download_excel_endpoint():
    if os.path.exists(EXCEL_PATH):
        return send_file(EXCEL_PATH, as_attachment=True, download_name="Inspection_Log.xlsx")
    elif os.path.exists(CSV_PATH):
        return send_file(CSV_PATH, as_attachment=True, download_name="Inspection_Log.csv")
    return "No inspection log generated yet", 404

# Authorization cache for scanned devices validated via PMS MQTT
validated_device_cache = {
    "serial_no": None,
    "apo_id": None,
    "timestamp": 0.0,
    "raw_input": None,
    "payloads": None,
    "takt_start_time": None
}

@app.route("/api/pms/validate", methods=["POST"])
def pms_validate_endpoint():
    """
    Validates a scanned device against PMS MQTT server before activating the START button:
    1. Extracts APoNumberId and DeviceSerialNo from scanned QR text (e.g. 'Id: ..; Sr: ..;').
    2. Sends DeviceAutoWhitelisting, SerialNoValidation, and StatusValidation sequentially.
    3. Confirms that the server responds with CmdType='ResponseFromServer', matching UtcRequest_ID,
       and RequestStatus='CONFIG SUCCESSFUL'.
    4. Caches the authorization and validated payloads so that the START button is activated
       and JSON payloads are populated.
    """
    global validated_device_cache
    data = request.get_json() or {}
    device_id = str(data.get("device_id") or "").strip()

    if not device_id:
        return jsonify({
            "status": "error",
            "validated": False,
            "message": "No serial number or QR code provided. Please scan a QR code first.",
            "payloads": build_empty_payloads(),
            "server_responses": {}
        }), 400

    apo_id, serial_no = parse_sticker(device_id)

    qr_scan_start_time = time.time()

    try:
        print("\n" + "=" * 60)
        print(f"[PMS VALIDATION GATE] Validating scanned device: Serial={serial_no}, APO={apo_id}...")
        print("=" * 60)

        res = pms_mqtt.run_startup_validation(
            device_serial_no=serial_no,
            apo_number_id=apo_id
        )

        validated_payloads = res.get("payloads") or {}
        server_responses = res.get("server_responses") or {}

        validated_device_cache = {
            "serial_no": serial_no,
            "apo_id": apo_id,
            "timestamp": time.time(),
            "raw_input": device_id,
            "payloads": validated_payloads,
            "server_responses": server_responses,
            "takt_start_time": qr_scan_start_time
        }

        status_was_checked = server_responses.get("StatusValidation") is not None

        if status_was_checked:
            print(
                f"[PMS VALIDATION GATE] Device, Serial, and Status validations "
                f"confirmed CONFIG SUCCESSFUL for {serial_no}. START button enabled.\n"
            )
        else:
            print(
                f"[PMS VALIDATION GATE] Device and Serial validations successful. "
                f"Status Validation skipped. START button enabled for {serial_no}.\n"
            )
        return jsonify({
            "status": "success",
            "validated": True,
            "serial_no": serial_no,
            "apo_id": apo_id,
            "raw_qr": device_id,
            "payloads": validated_payloads,
            "server_responses": server_responses,
            "message": (
                "Server responded CONFIG SUCCESSFUL for Whitelisting, SerialNo, and StatusValidation. START button activated."
                if server_responses.get("StatusValidation") is not None
                else "Whitelisting and SerialNo validations successful. StatusValidation skipped. START button activated."
            )
        })

    except Exception as exc:
        print(f"[PMS VALIDATION GATE] Validation failed for {serial_no}: {exc}\n")
        validated_device_cache = {
            "serial_no": None,
            "apo_id": None,
            "timestamp": 0.0,
            "raw_input": None,
            "payloads": None,
            "server_responses": None
        }
        return jsonify({
            "status": "error",
            "validated": False,
            "serial_no": serial_no,
            "apo_id": apo_id,
            "raw_qr": device_id,
            "message": str(exc),
            "payloads": build_empty_payloads(),
            "server_responses": {}
        }), 503

@app.route("/api/inspect", methods=["POST"])
def inspect_endpoint():
    global validated_device_cache
    t_start = time.time()
    data = request.get_json() or {}
    device_id = data.get("device_id") or "N/A"

    apo_id, serial_no = parse_sticker(device_id)

    # Pre-Inspection Validation Gate:
    # START inspection must ONLY proceed if DeviceAutoWhitelisting, SerialNoValidation,
    # and StatusValidation were sent and server responded CONFIG SUCCESSFUL matching UtcRequest_ID.
    now_t = time.time()
    is_prevalidated = (
        validated_device_cache.get("serial_no") == serial_no
        and validated_device_cache.get("payloads") is not None
        and (now_t - validated_device_cache.get("timestamp", 0.0)) < 600.0
    )

    if not is_prevalidated:
        print(f"[PMS GATE] Blocked unvalidated inspection attempt for Serial={serial_no}.")
        return jsonify({
            "status": "error",
            "stage": "pms_validation",
            "message": "Device not validated. Payloads must be sent and server must respond CONFIG SUCCESSFUL with matching UtcRequest_ID before START is allowed.",
            "inspection_started": False
        }), 403
    # Preserve the QR scan start time before clearing validation cache.
    takt_start_time = validated_device_cache.get("takt_start_time")

    # Reset cache after starting inspection so each board requires a fresh scan and validation
    validated_device_cache = {
        "serial_no": None,
        "apo_id": None,
        "timestamp": 0.0,
        "raw_input": None,
        "payloads": None,
        "takt_start_time": None
    }

    top_img_data = data.get("top_image")
    orig3_img_data = data.get("orig3_image") or data.get("orig2_image")

    response_payload = {}
    top_res = {}
    orig3_res = {}

    top_img_bgr = None
    orig3_img_bgr = None

    if top_img_data:
        top_img_bgr = decode_image(top_img_data)
        if top_img_bgr is not None:
            top_res = process_top_inspection(top_img_bgr)
            response_payload["top_wires_annotated"] = top_res["top_wires_annotated"]
            response_payload["top_origins_annotated"] = top_res["top_origins_annotated"]

    if orig3_img_data:
        orig3_img_bgr = decode_image(orig3_img_data)
        if orig3_img_bgr is not None:
            orig3_res = process_orig3_inspection(orig3_img_bgr)
            response_payload["orig3_annotated"] = orig3_res["orig3_annotated"]
            response_payload["orig2_annotated"] = orig3_res["orig3_annotated"]

    # Decision Engine: Top image detected by pose estimation model, side image by other model
    is_top_pass = top_res.get("top_pass", False)
    is_side_pass = orig3_res.get("side_pass", False)

    obl_pass = top_res.get("obl_pass", False)
    oum_pass = top_res.get("oum_pass", False)
    raw_obr_pass = top_res.get("obr_pass", False)
    c1_pass = top_res.get("c1_pass", False)
    c2_pass = top_res.get("c2_pass", False)
    c3_pass = top_res.get("c3_pass", False)
    orig3_grommet_pass = orig3_res.get("orig3_grommet_pass", False)

    effective_origin3_pass = is_side_pass
    wire_all_pass = is_top_pass
    origins_all_pass = is_top_pass and is_side_pass
    combined_pass = is_top_pass and is_side_pass

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    timestamp_file = datetime.now().strftime("%Y%m%d_%H%M%S")

    notes_lines = []
    notes_lines.append(f"═════════════ INSPECTION LOG: {serial_no} ═════════════")
    notes_lines.append(f"🏷️  Device Serial: {serial_no} | APO Number ID: {apo_id}\n")

    # 1. Top View (Pose Estimation Model)
    notes_lines.append("🔌 TOP VIEW (Arducam_TOP - Pose Estimation Model):")
    notes_lines.append(f" • Top Pose Inspection: {'✔ PASS' if is_top_pass else '✖ FAIL'}")
    top_lbls = top_res.get("detected_labels") or []
    if top_lbls:
        notes_lines.append(f" • Detections: {', '.join(top_lbls)}\n")
    else:
        notes_lines.append(" • Detections: None\n")

    # 2. Side View (Secondary Detection Model)
    notes_lines.append("🎯 SIDE VIEW (Arducam_ORIGIN - Secondary Detection Model):")
    notes_lines.append(f" • Side Detection Inspection: {'✔ PASS' if is_side_pass else '✖ FAIL'}")
    side_lbls = orig3_res.get("detected_labels") or []
    if side_lbls:
        notes_lines.append(f" • Detections: {', '.join(side_lbls)}\n")
    else:
        notes_lines.append(" • Detections: None\n")

    # Save snapshot images to Result/Pass or Result/Fail
    target_folder = "Pass" if combined_pass else "Fail"
    result_dir = os.path.join(os.getcwd(), "Result", target_folder)
    os.makedirs(result_dir, exist_ok=True)

    # Save annotated inspection views
    clean_device_id = "".join(c for c in serial_no if c.isalnum() or c in ('-', '_')).strip() or "device"
    prefix = f"{timestamp_file}_{clean_device_id}"

    if "img_wires" in top_res and top_res["img_wires"] is not None:
        cv2.imwrite(os.path.join(result_dir, f"{prefix}_top_wires.jpg"), top_res["img_wires"])
        cv2.imwrite(os.path.join(result_dir, "latest_top_wires.jpg"), top_res["img_wires"])

    if "img_origins" in top_res and top_res["img_origins"] is not None:
        cv2.imwrite(os.path.join(result_dir, f"{prefix}_top_origins.jpg"), top_res["img_origins"])
        cv2.imwrite(os.path.join(result_dir, "latest_top_origins.jpg"), top_res["img_origins"])

    if "img_o3" in orig3_res and orig3_res["img_o3"] is not None:
        cv2.imwrite(os.path.join(result_dir, f"{prefix}_orig3_grommet.jpg"), orig3_res["img_o3"])
        cv2.imwrite(os.path.join(result_dir, "latest_orig3_grommet.jpg"), orig3_res["img_o3"])

    # Save raw captures for traceability
    if top_img_bgr is not None:
        cv2.imwrite(os.path.join(result_dir, f"{prefix}_raw_top.jpg"), top_img_bgr)
    if orig3_img_bgr is not None:
        cv2.imwrite(os.path.join(result_dir, f"{prefix}_raw_orig3.jpg"), orig3_img_bgr)

    # Collect structured failure summary for Excel and UI
    # When the model runs and results into Pass then do nothing
    # and if it fails then in an fstring store what all cases have failed
    if combined_pass:
        # Pass -> do nothing
        failed_cases_str = ""
        notes_summary = "All checks passed"
        fail_details = []
    else:
        # Fail -> in an fstring store what all cases have failed
        failed_cases = []
        fail_details = list(top_res.get("fail_details") or [])

        if not is_side_pass:
            failed_cases.append("Origin 3 (Grommet)")
            fail_details.append({
                "item": "Origin 3 (Grommet)",
                "reason": "Origin 3 grommet not detected in side view"
            })
        if not obl_pass:
            failed_cases.append("Origin 1 (BL)")
        if not oum_pass:
            failed_cases.append("Origin 2 (UM)")
        if not raw_obr_pass:
            failed_cases.append("Origin 3 (BR Top)")
        if not c1_pass:
            failed_cases.append("Case 1 (BL)")
        if not c2_pass:
            failed_cases.append("Case 2 (UM)")
        if not c3_pass:
            failed_cases.append("Case 3 (BR)")

        seen = set()
        failed_cases_unique = [c for c in failed_cases if not (c in seen or seen.add(c))]
        failed_list = ", ".join(failed_cases_unique)
        failed_cases_str = f"Failed Cases: {failed_list}"
        notes_summary = failed_cases_str
        print(f"[MODEL INFERENCE] {failed_cases_str}")
        notes_lines.append(f"❌ {failed_cases_str}\n")

    fail_items = [f"{d['item']}: {d['reason']}" for d in fail_details]

    # 4. Log to Excel Sheet & CSV
    excel_record = {
        "timestamp": now_str,
        "device_id": serial_no,
        "combined_pass": combined_pass,
        "orig3_grommet_pass": is_side_pass,
        "obl_pass": obl_pass,
        "oum_pass": oum_pass,
        "obr_pass": raw_obr_pass,
        "c1_pass": c1_pass,
        "c2_pass": c2_pass,
        "c3_pass": c3_pass,
        "notes_summary": notes_summary,
        "failed_cases": "" if combined_pass else failed_cases_str,
        "saved_folder": f"Result/{target_folder}"
    }
    log_to_excel(excel_record)

    # 5. Generate and Log Selec PMS Report JSON Output (PASS or FAIL)
    # takt_time = time.time() - t_start
    # verdict_label = "PASS" if combined_pass else "FAIL"
    # Calculate TaktTime from QR scan until this point.

    # We'll finalize it after PMS accepts the ReportStore request.
    # Calculate elapsed time from QR scan.
    # takt_start_time was preserved before validated_device_cache was reset.
    # TaktTime: QR scan to AI inference final result.
    if takt_start_time is not None:
        takt_time = time.time() - takt_start_time
    else:
        takt_time = time.time() - t_start

    takt_time = round(takt_time, 3)

    verdict_label = "PASS" if combined_pass else "FAIL"
    try:
        req_id, pms_report_payload = build_report_store(
            device_serial_no=serial_no,
            apo_number_id=apo_id,
            takt_time=takt_time,
            vi_test_pass=combined_pass,
            remark=failed_cases_str,

            # Inspection timestamp in epoch milliseconds
            inspection_epoch_ms=int(time.time() * 1000)
        )

        # 1. Save in Result/PMS_Payloads/ with explicit PASS/FAIL prefix
        pms_json_path = save_payload_to_file(pms_report_payload, f"{verdict_label}_report_{clean_device_id}")

        # 2. Also save directly in Result/Pass/ or Result/Fail/ alongside the snapshot images
        folder_json_path = os.path.join(result_dir, f"{prefix}_{verdict_label}_report.json")
        with open(folder_json_path, "w", encoding="utf-8") as f:
            json.dump(pms_report_payload, f, indent=4)

        # 3. Save latest_pms_report.json in Result/Pass/ or Result/Fail/
        latest_folder_json = os.path.join(result_dir, "latest_pms_report.json")
        with open(latest_folder_json, "w", encoding="utf-8") as f:
            json.dump(pms_report_payload, f, indent=4)

        # 4. Save latest_run.json in Result/PMS_Payloads/
        latest_overall_json = os.path.join(os.path.dirname(pms_json_path), "latest_run.json")
        with open(latest_overall_json, "w", encoding="utf-8") as f:
            json.dump(pms_report_payload, f, indent=4)

        # response_payload["pms_report_json"] = pms_report_payload
        # response_payload["pms_json_file"] = folder_json_path
        # response_payload["verdict"] = verdict_label

        # # Publish the actual inspection result to PMS.(Aryan)
        # try:
        #     pms_response = pms_mqtt.publish_report(
        #         pms_report_payload
        #     )
        #     response_payload["pms_mqtt_response"] = pms_response

        # except Exception as exc:
        #     print(f"[PMS] ReportStore publish error: {exc}")
        #     response_payload["pms_mqtt_response"] = {
        #         "received": False,
        #         "message": str(exc)
        #     }

        
        response_payload["pms_report_json"] = pms_report_payload
        response_payload["pms_json_file"] = folder_json_path

        # Keep the AI inspection result separate from the final verdict.
        response_payload["ai_verdict"] = verdict_label

        # Publish ReportStore and wait for PMS validation.
        try:
            pms_response = pms_mqtt.publish_report(
                pms_report_payload
            )

            response_payload["pms_mqtt_response"] = pms_response

            # if (
            #     pms_response.get("received") is True
            #     and pms_response.get("accepted") is True
            # ):
            #     # PMS accepted ReportStore: use the AI inspection verdict.
            #     response_payload["verdict"] = verdict_label
            #     response_payload["pms_report_status"] = "ACCEPTED"

            if (
                pms_response.get("received") is True
                and pms_response.get("accepted") is True
            ):
                # PMS accepted ReportStore.
                # Calculate total elapsed time from QR scan to PMS acceptance.
                if takt_start_time is not None:
                    final_takt_time = time.time() - takt_start_time
                else:
                    final_takt_time = time.time() - t_start

                response_payload["takt_time_seconds"] = round(final_takt_time, 3)

                response_payload["verdict"] = verdict_label
                response_payload["pms_report_status"] = "ACCEPTED"

            elif pms_response.get("received") is True:
                # PMS responded but did not accept ReportStore.
                response_payload["verdict"] = "FAIL"
                response_payload["pms_report_status"] = "REJECTED"

                response_payload["pms_rejection_reason"] = (
                    pms_response.get("reason")
                    or pms_response.get("request_status")
                    or "PMS rejected the ReportStore request."
                )
                fail_details.append({
                    "item": "PMS Server (ReportStore)",
                    "reason": f"Rejected by server: {response_payload['pms_rejection_reason']}"
                })

            else:
                # No matching PMS response after the retry attempts.
                response_payload["verdict"] = "SERVER UNREACHABLE"
                response_payload["pms_report_status"] = "UNREACHABLE"

        except Exception as exc:
            print(f"[PMS] ReportStore publish error: {exc}")

            response_payload["pms_mqtt_response"] = {
                "received": False,
                "accepted": False,
                "message": str(exc)
            }


            response_payload["verdict"] = "SERVER UNREACHABLE"
            response_payload["pms_report_status"] = "UNREACHABLE"


        notes_lines.append(f"\n📡 PMS JSON REPORT ({verdict_label}):")
        notes_lines.append(f" • Verdict: {verdict_label} (ViTest: {1 if combined_pass else 0}, WireOrigin: {1 if origins_all_pass else 0})")
        notes_lines.append(f" • TaktTime: {round(takt_time, 3)}s")
        notes_lines.append(f" • Saved in Result/{target_folder}/: {os.path.basename(folder_json_path)}")
        notes_lines.append(f" • Saved in Result/PMS_Payloads/: {os.path.basename(pms_json_path)}")
    except Exception as e:
        print(f"[PMS Report] Warning: Could not generate PMS report: {e}")

    notes_lines.append(f"\n📊 Logged record to: Result/Inspection_Log.xlsx")
    notes_lines.append(f"📁 Saved result images to: Result/{target_folder}/")

    response_payload["results"] = {
        "obl_pass": obl_pass,
        "oum_pass": oum_pass,
        "obr_pass": is_side_pass,
        "raw_obr_pass": raw_obr_pass,
        "c1_pass": c1_pass,
        "c2_pass": c2_pass,
        "c3_pass": c3_pass,
        "orig3_grommet_pass": is_side_pass,
        "orig2_grommet_pass": is_side_pass
    }
    response_payload["fail_details"] = fail_details
    response_payload["fail_items"] = [f"{d['item']}: {d['reason']}" for d in fail_details]
    response_payload["failed_cases"] = failed_cases_str if not combined_pass else ""
    response_payload["failed_cases_str"] = failed_cases_str if not combined_pass else ""
    response_payload["is_top_pass"] = is_top_pass
    response_payload["is_side_pass"] = is_side_pass
    response_payload["device_id"] = serial_no
    response_payload["raw_qr"] = device_id
    response_payload["serial_no"] = serial_no
    response_payload["apo_id"] = apo_id
    response_payload["combined_pass"] = combined_pass
    response_payload["saved_folder"] = f"Result/{target_folder}"
    response_payload["excel_log"] = "Result/Inspection_Log.xlsx"
    response_payload["notes"] = "\n".join(notes_lines)

    return jsonify(response_payload)

@app.route("/api/pms/payloads", methods=["GET", "POST"])
def pms_payloads_endpoint():
    """
    Returns all 4 required PMS JSON payloads.
    The JSON fields will stay empty until the device has been scanned AND
    the server has responded with matching UtcRequest_ID and CONFIG SUCCESSFUL.
    """
    device_id = str(request.args.get("device_id") or "").strip()

    if not device_id:
        return jsonify({
            "status": "success",
            "detected_mac": "",
            "raw_qr": "",
            "scanned_serial_no": "",
            "device_serial_no": "",
            "apo_number_id": "",
            "payloads": build_empty_payloads(),
            "server_responses": {}
        })

    apo_id, serial_no = parse_sticker(device_id)

    # Check if this exact device has completed server validation with CONFIG SUCCESSFUL
    is_validated = (
        validated_device_cache.get("serial_no") == serial_no
        and validated_device_cache.get("payloads") is not None
        and (time.time() - validated_device_cache.get("timestamp", 0.0)) < 600.0
    )

    if not is_validated:
        # User requirement: "else these json fields will stay empty"
        return jsonify({
            "status": "unvalidated",
            "detected_mac": "",
            "raw_qr": device_id,
            "scanned_serial_no": serial_no,
            "device_serial_no": "",
            "apo_number_id": "",
            "payloads": build_empty_payloads(),
            "server_responses": {}
        })

    # Device is confirmed validated by server! Return the validated payloads.
    mac = get_pc_mac_address()
    validated_payloads = validated_device_cache.get("payloads") or {}
    server_responses = validated_device_cache.get("server_responses") or {}
    return jsonify({
        "status": "success",
        "detected_mac": mac,
        "raw_qr": device_id,
        "scanned_serial_no": serial_no,
        "device_serial_no": serial_no,
        "apo_number_id": apo_id,
        "payloads": validated_payloads,
        "server_responses": server_responses
    })

def open_browser():
    time.sleep(1.8)
    url = "http://127.0.0.1:7860/"
    opened = False
    try:
        opened = webbrowser.open(url, new=2)
    except Exception as e:
        print(f"[Browser] webbrowser.open: {e}")
    if not opened and sys.platform.startswith("win"):
        try:
            import subprocess
            subprocess.run(f'start "" "{url}"', shell=True)
        except Exception as e:
            print(f"[Browser] start command: {e}")

# if __name__ == "__main__":
#     print("\n" + "="*65)
#     print(" Dual-Camera Wiring & Origin Inspection System (Web App)")
#     print(" Running at: http://127.0.0.1:7860/")
#     print(" Inspection Log: Result/Inspection_Log.xlsx")
#     print(" Opening default web browser automatically...")
#     print("="*65 + "\n")
    
#     # Auto open browser in background thread
#     threading.Thread(target=open_browser, daemon=True).start()
#     app.run(host="127.0.0.1", port=7860, debug=False)

# ================= VERSION MANAGEMENT API =================

@app.route("/api/version", methods=["GET"])
def get_app_version():
    return jsonify({
        "success": True,
        "version": APP_VERSION
    })


@app.route("/api/production-line", methods=["GET"])
def get_production_line():
    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as config_file:
            config = json.load(config_file)

        production_line_no = str(
            config.get("plcJigProductionLineNo", "1")
        )

        return jsonify({
            "success": True,
            "production_line_no": production_line_no
        })

    except Exception as error:
        print(f"[GTI Version] Failed to read configuration: {error}")

        return jsonify({
            "success": False,
            "message": "Could not read GTI Version from pms_config.json."
        }), 500


@app.route("/api/admin/version", methods=["POST"])
def update_app_version():
    global APP_VERSION

    data = request.get_json(silent=True) or {}

    version = str(data.get("version", "")).strip()
    password = str(data.get("password", ""))

    # Validate admin password
    if password != ADMIN_VERSION_PASSWORD:
        return jsonify({
            "success": False,
            "message": "Incorrect admin password."
        }), 401

    # Validate version input
    if not version:
        return jsonify({
            "success": False,
            "message": "Please enter a version number."
        }), 400

    if len(version) > 30:
        return jsonify({
            "success": False,
            "message": "Version must be 30 characters or fewer."
        }), 400

    # Save version to pms_config.json under "slaveVersionNo".
    if not update_config_value("slaveVersionNo", version):
        return jsonify({
            "success": False,
            "message": "Could not save version to pms_config.json."
        }), 500

    # Update version in application memory
    APP_VERSION = version

    return jsonify({
        "success": True,
        "version": APP_VERSION,
        "message": "Version updated successfully."
    })


@app.route("/api/admin/jig-uid", methods=["POST"])
def update_admin_jig_uid():
    data = request.get_json(silent=True) or {}

    jig_uid = str(data.get("jig_uid", "")).strip()
    password = str(data.get("password", ""))

    # Validate admin password.
    if password != ADMIN_VERSION_PASSWORD:
        return jsonify({
            "success": False,
            "message": "Incorrect admin password."
        }), 401

    # Validate Jig UID.
    if not jig_uid:
        return jsonify({
            "success": False,
            "message": "Please enter a Jig UID."
        }), 400

    if len(jig_uid) > 100:
        return jsonify({
            "success": False,
            "message": "Jig UID must be 100 characters or fewer."
        }), 400

    # Save the new Jig UID without overwriting other settings.
    if not update_config_value("jigUid", jig_uid):
        return jsonify({
            "success": False,
            "message": "Could not save Jig UID to pms_config.json."
        }), 500

    return jsonify({
        "success": True,
        "jig_uid": jig_uid,
        "message": "Jig UID updated and saved successfully."
    })


@app.route("/api/admin/production-line", methods=["POST"])
def update_admin_production_line():
    data = request.get_json(silent=True) or {}

    production_line_no = str(
        data.get("production_line_no", "")
    ).strip()

    password = str(data.get("password", ""))

    # Validate admin password.
    if password != ADMIN_VERSION_PASSWORD:
        return jsonify({
            "success": False,
            "message": "Incorrect admin password."
        }), 401

    # Validate Production Line Number.
    if not production_line_no:
        return jsonify({
            "success": False,
            "message": "Please enter a GTI Version."
        }), 400

    if len(production_line_no) > 30:
        return jsonify({
            "success": False,
            "message": "GTI Version must be 30 characters or fewer."
        }), 400

    # Save the Production Line Number without overwriting other settings.
    if not update_config_value(
        "plcJigProductionLineNo",
        production_line_no
    ):
        return jsonify({
            "success": False,
            "message": "Could not save GTI Version to pms_config.json."
        }), 500

    return jsonify({
        "success": True,
        "production_line_no": production_line_no,
        "message": "GTI Version updated and saved successfully."
    })


@app.route("/api/config/validate", methods=["GET"])
def validate_pms_config():
    is_valid, error = validate_config_json()

    return jsonify({
        "success": True,
        "valid": is_valid,
        "error": error
    })





if __name__ == "__main__":
    print("\n" + "=" * 65)
    print(" Dual-Camera Wiring & Origin Inspection System (Web App)")
    print(" Running at: http://127.0.0.1:7860/")
    print(" Inspection Log: Result/Inspection_Log.xlsx")
    print(" Opening default web browser automatically...")
    print("=" * 65 + "\n")

    # Report pms_config.json problems at startup (the web page also shows a popup).
    _cfg_ok, _cfg_err = validate_config_json()
    if not _cfg_ok:
        print("[PMS Config] ERROR in pms_config.json:\n" + str(_cfg_err) + "\n")

    # Start MQTT connection in the background
    def start_mqtt_connection():
        try:
            print("[PMS MQTT] Starting connection...")
            pms_mqtt.connect()
            print("[PMS MQTT] Connected and subscribed successfully.")
        except Exception as e:
            print(f"[PMS MQTT] Connection failed: {e}")

    threading.Thread(
        target=start_mqtt_connection,
        daemon=True,
        name="PMS-MQTT-Connection"
    ).start()

    # Auto open browser in background thread
    threading.Thread(target=open_browser, daemon=True).start()

    app.run(host="127.0.0.1", port=7860, debug=False)